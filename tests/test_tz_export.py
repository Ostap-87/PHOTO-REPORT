"""Тесты разбора ТЗ (app/tz) и экспорта отчётов (app/export)."""

from __future__ import annotations

import io
import json
import os
import re
import sys
import zipfile

import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from app.export import (export_docx, export_pdf, export_registry_xlsx, export_xlsx, export_zip,  # noqa: E402
                        folder_name, folder_names)
from app.tz import parse_tz  # noqa: E402
from app.tz.common import extract_qty, make_title  # noqa: E402

SAMPLES = os.path.join(ROOT, "samples")


def _sample(name: str) -> bytes:
    path = os.path.join(SAMPLES, name)
    if not os.path.exists(path):
        sys.path.insert(0, SAMPLES)
        import make_samples  # noqa: PLC0415

        make_samples.main(SAMPLES)
    with open(path, "rb") as f:
        return f.read()


def _by_num(items: list[dict]) -> dict:
    return {it["num"]: it for it in items if it["num"]}


ALLOWED_TAGS = {"div", "p", "h1", "h2", "h3", "h4", "h5", "h6", "b", "i", "u", "br", "span", "table", "tbody",
                "tr", "td"}
ALLOWED_ATTRS = {"class", "data-src", "data-table", "style", "colspan", "rowspan"}


def _check_html_refs(res: dict) -> None:
    """Каждый item.source.ref есть в html; в html только безопасные теги и атрибуты."""
    from html.parser import HTMLParser

    html = res["html"]
    assert html
    refs: list[str] = []

    class Checker(HTMLParser):
        def handle_starttag(self, tag, attrs):
            assert tag in ALLOWED_TAGS, tag
            for name, value in attrs:
                assert name in ALLOWED_ATTRS, name
                if name == "data-src":
                    refs.append(value)
                if name == "style":
                    assert re.fullmatch(r"[a-z\-]+:[a-z0-9.]+(;[a-z\-]+:[a-z0-9.]+)*", value), value

    Checker().feed(html)
    assert len(refs) == len(set(refs)), "data-src должны быть уникальными"
    for it in res["items"]:
        assert it["source"]["ref"] in refs, it


# ---------------------------------------------------------------------------
# Разбор ТЗ
# ---------------------------------------------------------------------------

def test_banket_docx_word_numbering():
    res = parse_tz(_sample("tz_banket.docx"), "tz_banket.docx")
    items = res["items"]
    assert res["kind"] == "docx"
    assert "банкета" in res["title"].lower() and res["title"].startswith("Техническое задание")
    assert len(items) == 35
    assert {it["level"] for it in items} == {1, 2, 3}
    assert [it["num"] for it in items[:5]] == ["1", "1.1", "1.2", "1.3", "2"]
    nums = _by_num(items)
    assert nums["1"]["level"] == 1 and nums["1.1"]["level"] == 2 and nums["4.2.1"]["level"] == 3
    assert (nums["1.1"]["qty"], nums["1.1"]["unit"]) == (12, "шт")
    assert (nums["3.1"]["qty"], nums["3.1"]["unit"]) == ("6×4", "м")
    assert (nums["7.1"]["qty"], nums["7.1"]["unit"]) == (120, "перс")
    assert nums["4.1"]["qty"] is None  # «4 кВт» — параметр, а не количество
    assert nums["1.1"]["title"] == "Круглые банкетные столы диаметром 1,8 м — 12 шт"
    assert "Расстановка по схеме" in nums["1.1"]["text"]
    # маркированный список Word под нумерованным 7.1 — это текст пункта, а не отдельные пункты
    i71 = items.index(nums["7.1"])
    assert items[i71 + 1]["num"] == "7.2"
    assert nums["7.1"]["text"].split("\n")[1:] == [
        "• холодные закуски и салаты;", "• горячее блюдо на выбор гостя;", "• десерт и торт."]
    # шапка и подписи пунктами не стали
    joined = " ".join(it["text"] for it in items)
    assert "УТВЕРЖДАЮ" not in joined and "__________" not in joined and "Дата проведения" not in joined
    assert all(it["source"]["kind"] == "para" and it["source"]["ref"].startswith("p") for it in items)
    _check_html_refs(res)


def test_montazh_docx_table():
    res = parse_tz(_sample("tz_montazh.docx"), "tz_montazh.docx")
    items = res["items"]
    assert len(items) == 32
    assert {it["level"] for it in items} == {1, 2, 3}
    titles = [it["title"] for it in items]
    assert titles[0] == "Общие требования" and items[0]["level"] == 1
    assert [it["num"] for it in items[1:4]] == ["1", "2", "3"]  # стиль «List Number»
    rows = [it for it in items if it["source"]["kind"] == "row"]
    assert len(rows) == 20  # 4 раздела + 16 работ
    sections = [it for it in rows if not it["num"]]
    assert [s["title"] for s in sections] == ["Кабельные трассы", "Лотки", "Серверная", "Электрощиты"]
    assert all(s["level"] == 2 for s in sections)
    first = rows[1]
    assert (first["num"], first["level"], first["qty"], first["unit"]) == ("1", 3, 1850, "м")
    assert first["source"]["ref"] == "t0-r3"
    assert first["title"].startswith("Прокладка кабеля витая пара U/UTP кат. 6")
    units = {it["unit"] for it in rows if it["num"]}
    assert units == {"м", "шт", "компл", "линия"}
    assert [it["num"] for it in items[-2:]] == ["1)", "2)"]
    assert all(it["level"] == 2 for it in items[-2:])
    joined = " ".join(it["text"] for it in items)
    assert "Исполнитель:" not in joined and "Приложение" not in joined
    _check_html_refs(res)


def test_uborka_pdf():
    res = parse_tz(_sample("tz_uborka.pdf"), "tz_uborka.pdf")
    items = res["items"]
    assert res["kind"] == "pdf" and res["html"] is None
    assert len(res["pages"]) == 2 and all(p.startswith(b"\x89PNG") for p in res["pages"])
    assert "уборке" in res["title"]
    assert len(items) == 18
    assert {it["level"] for it in items} == {1, 2}
    nums = _by_num(items)
    assert (nums["1.1"]["qty"], nums["1.1"]["unit"]) == (850, "м²")
    # длинный пункт на нескольких строках, перенос начинается с тире — это один пункт
    assert (nums["3.1"]["qty"], nums["3.1"]["unit"]) == (210, "м²")
    assert nums["3.1"]["title"] == "Мытьё окон изнутри и снаружи"
    assert "охраны труда" in nums["3.1"]["text"]
    assert nums["а)"]["level"] == 2 and nums["б)"]["level"] == 2
    assert "• Персонал в униформе" in nums["4"]["text"]  # маркеры под «4.» — текст пункта
    for it in items:
        src = it["source"]
        assert src["ref"] is None and src["page"] in (1, 2)
        assert len(src["bbox"]) == 4 and all(0 <= v <= 1 for v in src["bbox"])
    assert nums["4"]["source"]["page"] == 2
    joined = " ".join(it["text"] for it in items)
    assert "Страница" not in joined and "УТВЕРЖДАЮ" not in joined and "Исполнитель:" not in joined


def test_vystavka_xlsx():
    res = parse_tz(_sample("tz_vystavka.xlsx"), "tz_vystavka.xlsx")
    items = res["items"]
    assert res["kind"] == "xlsx"
    assert res["title"] == "Техническое задание на застройку выставочного стенда"
    assert len(items) == 22
    assert {it["level"] for it in items} == {1, 2}
    nums = _by_num(items)
    assert nums["1"]["level"] == 1 and nums["1"]["title"] == "Конструкция стенда"
    assert (nums["1.1"]["qty"], nums["1.1"]["unit"]) == (36, "м²")
    assert (nums["1.2"]["qty"], nums["1.2"]["unit"]) == (12, "п.м")
    assert (nums["5.2"]["qty"], nums["5.2"]["unit"]) == (1, "усл")
    assert nums["1.1"]["title"] == "Подиум высотой 100 мм с ковровым покрытием"  # описание — в text, не в title
    assert "серый ковролин" in nums["1.1"]["text"]
    assert not any("Итого" in it["title"] for it in items)
    _check_html_refs(res)


def test_forum_table_docx_cells():
    """ТЗ одной таблицей «№ | Раздел | Требования»: пункты из абзацев ячейки, автонумерация Word."""
    res = parse_tz(_sample("tz_forum_table.docx"), "tz_forum_table.docx")
    items = res["items"]
    assert [it["num"] for it in items] == [
        "1", "1.1", "2", "2.1", "2.2", "3", "3.1", "3.2", "4", "4.1", "4.2",
        "5", "5.1", "5.1.1", "5.1.2", "5.1.2.1", "5.2", "5.2.1", "5.2.2", "5.3", "5.3.1", "5.4",
        "6", "6.1", "6.1.1", "6.1.2", "6.1.3", "6.1.4", "6.1.5", "6.2", "7"]
    assert all(it["level"] == it["num"].count(".") + 1 for it in items)
    by = _by_num(items)
    assert by["5"]["title"] == "Требования к оказываемым услугам" and by["5"]["source"]["ref"] == "t0-r5"
    # маркеры «-», абзацы без номера и локальный список «1.», «2.» — текст пункта, а не пункты
    assert by["5.1.2"]["text"].split("\n") == [
        "Оснащение главного зала:", "- не менее 300 посадочных мест;", "- сцена размером не менее 8×5 м;",
        "- светодиодный экран 6×3 м;", "- акустическая система — не менее 8 шт.",
        "Перечень оборудования согласовывается с Заказчиком."]
    assert by["5.1.2.1"]["text"].split("\n")[1:] == [
        "1. Резервный ноутбук с презентациями.", "2. Запасные радиомикрофоны — 2 шт."]
    assert by["5.3"]["text"].split("\n")[1:] == ["• на всех экранах площадки;", "• в окне онлайн-трансляции."]
    # ручной перенос строки (w:br) сохраняется
    assert by["5.3.1"]["text"] == "Подтверждение участия отправляется на почту.\nПовторное письмо — за сутки до Форума."
    assert (by["6.1.2"]["qty"], by["6.1.2"]["unit"]) == (15, "шт")
    assert by["6.1.3"]["title"].startswith("Фотографии, сгруппированные по каждому мероприятию")
    assert by["7"]["text"].count("\n") == 2  # раздел без номеров в ячейке — один пункт с текстом
    src = by["6.1.2"]["source"]
    assert (src["kind"], src["ref"], src["row"], src["cell"]) == ("cell", "t0-r6-p4", 6, 2)
    _check_html_refs(res)
    assert 'data-src="t0-r5-p19"' in res["html"] and "<br>" in res["html"]
    ft = res["full_text"]
    assert "\n5. | Требования к оказываемым услугам | 5.1. Исполнитель обеспечивает площадку" in ft
    for line in ("3.1. Услуги оказываются", "5.2.1. Регистрация ведётся", "6.1.2. Скриншоты в количестве 15 штук",
                 "6.1.5. Акт оказанных услуг.", "6.2. Отчёт согласовывается", "- на всех экранах площадки;",
                 "- предпринимателей;", "Повторное письмо — за сутки до Форума."):
        assert "\n" + line in ft or "| " + line in ft, line  # абзац с новой строки или первый в ячейке
    assert ft.index("5.1.2.1.") < ft.index("\n5.2. ") < ft.index("\n6.1.1.")


def test_full_text_formats():
    ft = parse_tz(_sample("tz_banket.docx"), "tz_banket.docx")["full_text"]
    lines = ft.split("\n")
    assert lines[0] == "УТВЕРЖДАЮ"
    assert "1. Мебель и расстановка" in lines and "4.2.1. Комплект запасных батарей на весь период мероприятия." in lines
    assert "- холодные закуски и салаты;" in lines  # маркированный список Word
    ft = parse_tz(_sample("tz_montazh.docx"), "tz_montazh.docx")["full_text"]
    assert "№ п/п | Наименование работ | Ед. изм. | Кол-во" in ft and "\nКабельные трассы\n" in ft
    assert "1 | Прокладка кабеля витая пара U/UTP кат. 6 в гофрированной трубе | м | 1 850" in ft
    ft = parse_tz(_sample("tz_uborka.pdf"), "tz_uborka.pdf")["full_text"]
    assert ft.startswith("--- стр. 1 ---\nУТВЕРЖДАЮ") and "\n--- стр. 2 ---\n" in ft
    assert ft.index("1.1. Влажная уборка") < ft.index("--- стр. 2 ---") < ft.index("4. Требования к исполнителю")
    ft = parse_tz(_sample("tz_vystavka.xlsx"), "tz_vystavka.xlsx")["full_text"]
    assert "№ п/п | Наименование | Описание / требования | Ед. изм. | Кол-во" in ft
    assert "1.1 | Подиум высотой 100 мм с ковровым покрытием | серый ковролин, алюминиевая окантовка | м² | 36" in ft
    assert "--- лист «Контакты» ---\nКонтактное лицо | Телефон | Почта" in ft
    ft = parse_tz("Позиция;Наименование;Ед.;Кол-во\n1;Стол;шт;10\n".encode("cp1251"), "t.csv")["full_text"]
    assert ft == "Позиция | Наименование | Ед. | Кол-во\n1 | Стол | шт | 10"
    txt = "Техническое задание\r\n1. Пункт один\r\n- деталь  \r\n"
    assert parse_tz(txt.encode(), "t.txt")["full_text"] == "Техническое задание\n1. Пункт один\n- деталь"


def test_text_csv_markdown():
    txt = ("УТВЕРЖДАЮ\nДиректор\nТЕХНИЧЕСКОЕ ЗАДАНИЕ\nна уборку территории\n"
           "1. Уборка снега — 1 200 кв. м\nежедневно до 8:00.\n2. Вывоз снега\n2.1 Погрузка — 50 м3\n"
           "2.2 Утилизация\n- талоны\n- акт\nЗаказчик: ________ / Иванов И.И.\n")
    res = parse_tz(txt.encode("cp1251"), "tz.txt")
    assert res["kind"] == "txt" and res["title"] == "Техническое задание на уборку территории"
    got = [(it["num"], it["level"]) for it in res["items"]]
    assert got == [("1", 1), ("2", 1), ("2.1", 2), ("2.2", 2)]
    assert res["items"][3]["text"] == "Утилизация\n- талоны\n- акт"
    assert (res["items"][0]["qty"], res["items"][0]["unit"]) == (1200, "м²")
    assert "ежедневно" in res["items"][0]["text"]
    _check_html_refs(res)

    md = ("# Техническое задание на вывеску\n\n## Работы\n1. Изготовление короба 3×1 м\n   - фрезеровка\n"
          "2. Монтаж — 1 шт.\n\n| № | Наименование работ | Кол-во | Ед. изм. |\n|---|---|---|---|\n"
          "| 1 | Подключение подсветки | 1 | компл. |\n\n**Подписи сторон**\nЗаказчик ____\n")
    res = parse_tz(md.encode(), "tz.md")
    got = [(it["num"], it["level"], it["title"]) for it in res["items"]]
    assert got[0] == ("", 1, "Работы")
    assert ("1", 2, "Изготовление короба 3×1 м") in got
    assert "- фрезеровка" in res["items"][1]["text"]
    assert res["items"][-1]["source"]["ref"] == "t0-r1" and res["items"][-1]["unit"] == "компл"
    _check_html_refs(res)

    csv_data = "Позиция;Наименование;Ед.;Кол-во\n1;Стол;шт;10\n2;Стул;шт.;40\n;Итого;;50\n".encode("cp1251")
    res = parse_tz(csv_data, "tz.csv")
    assert res["kind"] == "csv"
    assert [(it["num"], it["title"], it["qty"], it["unit"]) for it in res["items"]] == [
        ("1", "Стол", 10, "шт"), ("2", "Стул", 40, "шт")]


def test_docx_escaping_and_table_autonumber():
    from docx import Document
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls, qn

    d = Document()
    d.add_paragraph("Техническое задание на клининг")
    p = d.add_paragraph()
    p.add_run("Офисная зона").bold = True
    d.add_paragraph("Мытьё полов — 300 м²")
    d.add_paragraph("Протирка столов — 40 шт.")
    numbering = d.part.numbering_part.element
    aid = [n for n in numbering.findall(qn("w:num")) if n.get(qn("w:numId")) == "5"][0] \
        .find(qn("w:abstractNumId")).get(qn("w:val"))
    numbering.append(parse_xml(f'<w:num {nsdecls("w")} w:numId="78"><w:abstractNumId w:val="{aid}"/>'
                               f'<w:lvlOverride w:ilvl="0"><w:startOverride w:val="1"/></w:lvlOverride></w:num>'))
    t = d.add_table(rows=3, cols=3)
    for c, h in enumerate(("№", "Наименование", "Количество")):
        t.cell(0, c).text = h
    for r, (name, q) in enumerate([("Мойка окон", "120 м2"), ("Химчистка ковров", "по факту")], start=1):
        np_ = t.cell(r, 0).paragraphs[0]._p.get_or_add_pPr().get_or_add_numPr()
        np_.get_or_add_numId().val = 78
        np_.get_or_add_ilvl().val = 0
        t.cell(r, 1).text = name
        t.cell(r, 2).text = q
    d.add_paragraph('<script>alert(1)</script> & <img src=x onerror="alert(1)">')
    buf = io.BytesIO()
    d.save(buf)
    res = parse_tz(buf.getvalue(), "x.docx")
    got = [(it["num"], it["level"], it["title"], it["qty"], it["unit"]) for it in res["items"]]
    assert got[:3] == [("", 1, "Офисная зона", None, ""), ("", 2, "Мытьё полов — 300 м²", 300, "м²"),
                       ("", 2, "Протирка столов — 40 шт", 40, "шт")]
    assert got[3][:2] == ("1", 2) and got[3][3:] == (120, "м²")  # автонумерация Word в колонке «№»
    assert got[4][:2] == ("2", 2) and got[4][3] == "по факту"
    assert "&lt;script&gt;" in res["html"] and "&lt;img" in res["html"]
    _check_html_refs(res)


def test_scan_pdf_and_errors():
    import pymupdf
    from PIL import Image

    doc = pymupdf.open()
    page = doc.new_page()
    b = io.BytesIO()
    Image.new("RGB", (300, 400), "white").save(b, "PNG")
    page.insert_image(page.rect, stream=b.getvalue())
    res = parse_tz(doc.tobytes(), "скан.pdf")
    assert res["items"] == [] and len(res["pages"]) == 1 and res["warnings"]
    with pytest.raises(ValueError):
        parse_tz(b"\xd0\xcf\x11\xe0" + b"\x00" * 100, "old.doc")
    with pytest.raises(ValueError):
        parse_tz(b"", "empty.txt")
    with pytest.raises(ValueError):
        parse_tz(b"not a docx", "broken.docx")


@pytest.mark.parametrize("text,expected", [
    ("Стулья — 120 шт.", (120, "шт")),
    ("Мытьё окон — 210 кв. м", (210, "м²")),
    ("Кабель 1 850 м", (1850, "м")),
    ("Подиум 6×4 м, высота 0,4 м", ("6×4", "м")),
    ("Отверстия диаметром до 50 мм — 36 шт", (36, "шт")),
    ("Запас кабеля не менее 1,5 м на линию", (None, "")),
    ("Звук мощностью 4 кВт", (None, "")),
    ("Демонтаж в течение 3 часов", (None, "")),
    ("Плитка 45,5 м2", (45.5, "м²")),
])
def test_extract_qty(text, expected):
    assert extract_qty(text) == expected


def test_make_title():
    assert make_title("Круглые столы — 12 шт. Расстановка по схеме.") == "Круглые столы — 12 шт"
    long = "Мытьё окон изнутри и снаружи, " + "включая рамы и откосы " * 10
    assert make_title(long) == "Мытьё окон изнутри и снаружи"
    assert len(make_title("а" * 300)) <= 121
    assert make_title("Доставка в г. Москва, ул. Мира. Далее текст") == "Доставка в г. Москва, ул. Мира"


# ---------------------------------------------------------------------------
# Экспорт
# ---------------------------------------------------------------------------

def _jpeg(path: str, color: tuple, size=(1200, 900)) -> None:
    from PIL import Image, ImageDraw

    im = Image.new("RGB", size, color)
    d = ImageDraw.Draw(im)
    d.rectangle([size[0] // 4, size[1] // 4, size[0] * 3 // 4, size[1] * 3 // 4], fill=(255, 255, 255))
    os.makedirs(os.path.dirname(path), exist_ok=True)
    im.save(path, "JPEG", quality=85)


@pytest.fixture()
def project(tmp_path):
    pdir = tmp_path / "p1"
    for sub in ("tz", "photos/original", "photos/work", "photos/thumb"):
        (pdir / sub).mkdir(parents=True, exist_ok=True)
    items = [
        {"id": "it1", "num": "1", "level": 1, "title": "Мебель", "text": "Мебель", "qty": None, "unit": "",
         "source": {"kind": "para", "index": 1, "page": None, "ref": "p1"}, "status": "todo", "comment": "",
         "done_at": "", "hint": "", "photo_ids": []},
        {"id": "it2", "num": "1.1", "level": 2, "title": "Столы круглые — 12 шт",
         "text": "Столы круглые — 12 шт. Расстановка по схеме.", "qty": 12, "unit": "шт",
         "source": {"kind": "para", "index": 2, "page": None, "ref": "p2"}, "status": "done",
         "comment": "=HYPERLINK(\"http://example.com\")", "done_at": "2026-10-21", "hint": "",
         "photo_ids": ["ph1", "ph2", "ph_missing"]},
        {"id": "it3", "num": "1.2", "level": 2, "title": "Стулья: Кьявари / белые?", "text": "Стулья", "qty": 120,
         "unit": "шт", "source": {"kind": "para", "index": 3, "page": None, "ref": "p3"}, "status": "provided",
         "comment": "", "done_at": "2026-10-22", "hint": "", "photo_ids": ["ph3"]},
        {"id": "it4", "num": "2", "level": 1, "title": "Сцена 6×4 м", "text": "Сцена 6×4 м", "qty": "6×4",
         "unit": "м", "source": {"kind": "row", "index": 3, "table": 0, "row": 3, "page": None, "ref": "t0-r3"},
         "status": "executed", "comment": "+ссылка", "done_at": "", "hint": "", "photo_ids": ["ph4"]},
        {"id": "it5", "num": "3", "level": 1, "title": "Свет", "text": "Свет", "qty": None, "unit": "",
         "source": {"kind": "para", "index": 5, "page": None, "ref": "p5"}, "status": "todo", "comment": "",
         "done_at": "", "hint": "", "photo_ids": []},
    ]

    def photo(pid, name, item_id, **kw):
        base = {"id": pid, "name": name, "ext": os.path.splitext(name)[1].lower(), "sha256": "", "w": 1200, "h": 900,
                "exif": {"datetime": "2026-10-21 14:03:11", "gps": [55.75, 37.62], "camera": "Тестовая камера"},
                "quality": {"brightness": 0.3, "flags": ["dark"], "needs_enhance": True, "reshoot": False},
                "enhanced": True, "applied": ["exposure"], "faces": 1, "faces_blurred": True, "item_id": item_id,
                "confidence": 0.81, "suggestions": [[item_id, 0.81]], "confirmed": True, "duplicate_of": None,
                "status": "ready", "error": ""}
        base.update(kw)
        return base

    photos = [photo("ph1", "IMG_0001.JPG", "it2"), photo("ph2", "IMG_0002.jpg", "it2", enhanced=False),
              photo("ph3", "IMG_0003.png", "it3"), photo("ph4", "IMG_0004.jpg", "it4"),
              photo("ph5", "=cmd.jpg", None, confidence=0.2, confirmed=False),
              photo("ph_missing", "IMG_9999.jpg", "it2")]
    for pid, col in (("ph1", (200, 60, 60)), ("ph2", (60, 160, 60)), ("ph4", (60, 60, 200)), ("ph5", (120, 120, 0))):
        _jpeg(str(pdir / "photos" / "work" / f"{pid}.jpg"), col)
        _jpeg(str(pdir / "photos" / "original" / f"{pid}.jpg"), col)
    from PIL import Image  # у ph3 нет рабочей копии — берётся исходник PNG

    Image.new("RGB", (900, 1200), (90, 90, 90)).save(str(pdir / "photos" / "original" / "ph3.png"))
    proj = {"id": "p1", "name": "Тестовый банкет", "customer": "ООО «Пример»", "created": "2026-10-20T10:00:00",
            "updated": "2026-10-22T10:00:00", "settings": {"enhance": "auto", "blur_faces": True, "face_mode": "blur"},
            "tz": {"file": "tz/tz.docx", "kind": "docx", "title": "Техническое задание на банкет",
                   "html": "tz/source.html", "pages": 0},
            "items": items, "photos": photos}
    (pdir / "project.json").write_text(json.dumps(proj, ensure_ascii=False), encoding="utf-8")
    return json.loads((pdir / "project.json").read_text(encoding="utf-8")), str(pdir)


def test_export_docx(project, tmp_path):
    from docx import Document

    proj, pdir = project
    out = export_docx(proj, pdir, str(tmp_path / "out" / "отчёт.docx"), {"include_unsorted": True})
    doc = Document(out)
    assert len(doc.inline_shapes) == 5  # 4 фото в пунктах (одно без файла) + 1 неразобранное
    text = "\n".join(p.text for p in doc.paragraphs)
    for t in ("Тестовый банкет", "Фото без пункта ТЗ", "Исполнение по пунктам ТЗ", "Комментарий:"):
        assert t in text
    cells = " ".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert "Фото 1 · 21.10.2026 14:03 · улучшено · лица скрыты" in cells
    assert "Фото 2 · 21.10.2026 14:03 · лица скрыты" in cells
    assert "Выполнено" in cells and "Предоставлено" in cells and "Исполнено" in cells
    # без «Неразобранного» и только выполненные
    out2 = export_docx(proj, pdir, str(tmp_path / "only.docx"), {"only_status": ["done"], "photos_per_row": 3})
    assert len(Document(out2).inline_shapes) == 2
    # картинки уменьшены до 1600 px и сжаты
    with zipfile.ZipFile(out) as z:
        media = [n for n in z.namelist() if n.startswith("word/media/")]
        assert len(media) == 5


def test_export_pdf(project, tmp_path):
    import pymupdf

    proj, pdir = project
    out = export_pdf(proj, pdir, str(tmp_path / "отчёт.pdf"), {"include_unsorted": True, "photos_per_row": 2})
    doc = pymupdf.open(out)
    text = "".join(page.get_text() for page in doc)
    assert "Тестовый банкет" in text and "Выполнено" in text and "Фото без пункта ТЗ" in text
    assert f"стр. 1 из {doc.page_count}" in text
    images = sum(len(page.get_images()) for page in doc)
    assert images == 5
    assert doc[0].rect.width < doc[0].rect.height  # A4 книжная


def test_export_xlsx(project, tmp_path):
    import openpyxl

    proj, pdir = project
    out = export_xlsx(proj, pdir, str(tmp_path / "отчёт.xlsx"))
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Исполнение", "Фото"]
    ws = wb["Исполнение"]
    assert [c.value for c in ws[1]] == ["№", "Пункт ТЗ", "Кол-во", "Ед.", "Статус", "Дата", "Комментарий",
                                         "Фото, шт.", "Файлы фото"]  # без новых полей — прежние колонки
    row = [c.value for c in ws[3]]
    assert row[0] == "1.1" and row[2] == 12 and row[4] == "Выполнено" and row[7] == 3
    assert row[8] == "IMG_0001.JPG; IMG_0002.jpg; IMG_9999.jpg"
    comment = ws.cell(row=3, column=7)
    assert comment.value.startswith("=HYPERLINK") and comment.data_type == "s"  # не формула
    assert ws.cell(row=5, column=7).data_type == "s"
    ws2 = wb["Фото"]
    assert [c.value for c in ws2[1]] == ["Файл", "Пункт ТЗ", "Дата съёмки", "GPS", "Камера", "Качество",
                                          "Улучшено", "Лица скрыты", "Уверенность автосортировки", "Подтверждено"]
    names = [ws2.cell(row=r, column=1).value for r in range(2, ws2.max_row + 1)]
    assert len(names) == 6 and "=cmd.jpg" in names
    assert ws2.cell(row=2 + names.index("=cmd.jpg"), column=1).data_type == "s"
    r1 = [c.value for c in ws2[2]]
    assert r1[1] == "1.1 Столы круглые — 12 шт" and r1[3] == "55.750000, 37.620000" and r1[6] == "Да"
    with zipfile.ZipFile(out) as z:
        sheet_xml = z.read("xl/worksheets/sheet1.xml").decode()
        assert "<f>" not in sheet_xml


def test_export_zip(project, tmp_path):
    proj, pdir = project
    out = export_zip(proj, pdir, str(tmp_path / "фото.zip"), "work")
    names = zipfile.ZipFile(out).namelist()
    fn = folder_names(proj)
    assert fn["it2"] == "02 1.1 Столы круглые — 12 шт" == folder_name(proj["items"][1], 2)
    assert folder_name(proj["items"][1]) == "1.1 Столы круглые — 12 шт"
    assert f"{fn['it2']}/1.1_01_IMG_0001.jpg" in names
    assert f"{fn['it2']}/1.1_02_IMG_0002.jpg" in names
    assert fn["it3"] == "03 1.2 Стулья Кьявари белые"  # недопустимые символы убраны
    assert f"{fn['it3']}/1.2_01_IMG_0003.png" in names
    assert any(n.startswith("Неразобранное/001_") for n in names)
    assert not any(n.startswith(fn["it5"]) for n in names)  # пустые пункты папок не дают
    assert len(names) == 5
    assert all(len(n.split("/")[0]) <= 80 for n in names)
    out2 = export_zip(proj, pdir, str(tmp_path / "orig.zip"), "original")
    assert f"{fn['it2']}/1.1_01_IMG_0001.jpg" in zipfile.ZipFile(out2).namelist()
    long_item = {"num": "12.3.4", "title": "Очень длинное название пункта " * 10}
    assert len(folder_name(long_item, 7)) <= 80 and folder_name(long_item, 7).startswith("07 12.3.4 Очень")


def test_export_from_parsed_tz(tmp_path):
    """Сквозной сценарий: пункты из образца ТЗ → проект → все форматы отчёта."""
    res = parse_tz(_sample("tz_vystavka.xlsx"), "tz_vystavka.xlsx")
    pdir = tmp_path / "p2"
    items = []
    for i, it in enumerate(res["items"], start=1):
        items.append({"id": f"it{i}", **it, "status": "done" if i % 2 else "todo", "comment": "", "done_at": "",
                      "hint": "", "photo_ids": []})
    _jpeg(str(pdir / "photos" / "work" / "a.jpg"), (10, 120, 200))
    items[1]["photo_ids"] = ["a"]
    proj = {"id": "p2", "name": "Стенд", "customer": "", "items": items,
            "photos": [{"id": "a", "name": "a.jpg", "ext": ".jpg", "item_id": items[1]["id"], "exif": {}}],
            "tz": {"title": res["title"]}}
    for fn, f in (("r.docx", export_docx), ("r.pdf", export_pdf)):
        assert os.path.getsize(f(proj, str(pdir), str(tmp_path / fn), {})) > 1000
    assert os.path.getsize(export_xlsx(proj, str(pdir), str(tmp_path / "r.xlsx"))) > 1000
    assert len(zipfile.ZipFile(export_zip(proj, str(pdir), str(tmp_path / "r.zip"))).namelist()) == 1


# ---------------------------------------------------------------------------
# Необязательные поля анализа ТЗ: раздел, «что приложить», комментарий, ссылки, файлы, реестр
# ---------------------------------------------------------------------------

SAFE_URL = "https://example.com/post?a=1&b=2"


@pytest.fixture()
def rich_project(project):
    proj, pdir = project
    items = {it["id"]: it for it in proj["items"]}
    items["it2"].update(kind="registry", section="6. Требования к услугам",
                        required="Фото расстановки столов, не менее 3 шт., JPG",
                        note="=Уточнить у заказчика схему", links=[SAFE_URL, "javascript:alert(1)"])
    items["it4"].update(kind="registry", num="6.8–6.12, 7.1.7", section="6. Требования к услугам",
                        required="Скриншоты публикаций, PDF", note="")
    items["it5"].update(kind="tz")
    files_dir = os.path.join(pdir, "photos", "files")
    os.makedirs(files_dir, exist_ok=True)
    for fid, ext, data in (("f1", ".pdf", b"%PDF-1.4 test"), ("f2", ".mp4", b"\x00" * 64),
                           ("f3", ".xlsx", b"PK test")):
        with open(os.path.join(files_dir, fid + ext), "wb") as f:
            f.write(data)
    proj["files"] = [
        {"id": "f1", "name": "Акт приёмки.pdf", "ext": ".pdf", "size": 120_000, "item_id": "it2", "added": "1"},
        {"id": "f2", "name": "Видео сцены.mp4", "ext": ".mp4", "size": 35_000_000, "item_id": "it5", "added": "2"},
        {"id": "f3", "name": "Список гостей.xlsx", "ext": ".xlsx", "size": 900, "item_id": None, "added": "3"},
        {"id": "f_missing", "name": "Нет на диске.docx", "ext": ".docx", "size": 10, "item_id": "it2"},
    ]
    proj["clarify"] = [{"num": "6.40", "section": "6. Требования к услугам", "name": "Брендирование",
                        "description": "Объём брендирования не указан", "reason": "+нет количества"}]
    return proj, pdir


def test_export_extra_fields_docx_pdf(rich_project, tmp_path):
    from docx import Document
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    import pymupdf

    proj, pdir = rich_project
    out = export_docx(proj, pdir, str(tmp_path / "r.docx"), {"include_unsorted": True})
    doc = Document(out)
    text = "\n".join(p.text for p in doc.paragraphs)
    for t in ("Раздел: 6. Требования к услугам", "Что приложить: Фото расстановки столов, не менее 3 шт., JPG",
              "Комментарий анализа: =Уточнить у заказчика схему", "Ссылки:", SAFE_URL, "Файлы:",
              "Акт приёмки.pdf (117 КБ)", "Видео сцены.mp4 (33,4 МБ)", "Файлы без пункта ТЗ:", "Список гостей.xlsx",
              "6.8–6.12, 7.1.7 Сцена 6×4 м"):
        assert t in text, t
    # порядок карточки: Раздел → Что приложить → комментарий анализа → статус → ссылки → файлы
    i = text.index("Раздел: 6. Требования")
    assert i < text.index("Что приложить: Фото") < text.index("Комментарий анализа") < text.index("Ссылки:") \
        < text.index("Акт приёмки.pdf")
    links = [r.target_ref for r in doc.part.rels.values() if r.reltype == RT.HYPERLINK]
    assert links == [SAFE_URL]  # javascript: — только текстом
    assert len(doc.inline_shapes) == 5

    out = export_pdf(proj, pdir, str(tmp_path / "r.pdf"), {"include_unsorted": True})
    pdf = pymupdf.open(out)
    ptext = "".join(pg.get_text() for pg in pdf)
    for t in ("Раздел:", "Что приложить:", "Комментарий анализа:", "Акт приёмки.pdf", "Видео сцены.mp4",
              "Список гостей.xlsx", "javascript:alert(1)"):
        assert t in ptext, t
    uris = [ln.get("uri") for pg in pdf for ln in pg.get_links() if ln.get("uri")]
    assert uris == [SAFE_URL]


def test_export_extra_fields_xlsx_zip(rich_project, tmp_path):
    import openpyxl

    proj, pdir = rich_project
    wb = openpyxl.load_workbook(export_xlsx(proj, pdir, str(tmp_path / "r.xlsx")))
    ws = wb["Исполнение"]
    head = [c.value for c in ws[1]]
    assert head == ["№", "Раздел", "Пункт ТЗ", "Что приложить", "Кол-во", "Ед.", "Статус", "Дата", "Комментарий",
                    "Комментарий анализа", "Ссылки", "Фото, шт.", "Файлов", "Файлы фото"]
    row = {h: ws.cell(row=3, column=i + 1) for i, h in enumerate(head)}
    assert row["Раздел"].value == "6. Требования к услугам"
    assert row["Что приложить"].value.startswith("Фото расстановки")
    assert row["Комментарий анализа"].value == "=Уточнить у заказчика схему"
    assert row["Комментарий анализа"].data_type == "s"
    assert row["Ссылки"].value == SAFE_URL + "\njavascript:alert(1)"
    assert row["Файлов"].value == 2 and row["Фото, шт."].value == 3

    names = zipfile.ZipFile(export_zip(proj, pdir, str(tmp_path / "r.zip"))).namelist()
    fn = folder_names(proj)
    assert f"{fn['it2']}/Акт приёмки.pdf" in names
    assert f"{fn['it5']}/Видео сцены.mp4" in names  # папка появилась из-за файла, фото в пункте нет
    assert "Неразобранное/Список гостей.xlsx" in names
    assert not any("Нет на диске" in n for n in names)
    assert fn["it4"].startswith("04 6.8–6.12, 7.1.7")


def test_export_registry_xlsx(rich_project, tmp_path):
    import openpyxl

    proj, _pdir = rich_project
    out = export_registry_xlsx(proj, str(tmp_path / "реестр.xlsx"))
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Реестр подтверждений", "Требуют уточнения"]
    ws = wb["Реестр подтверждений"]
    assert [c.value for c in ws[1]] == ["Пункт ТЗ", "Раздел", "Наименование", "Описание",
                                         "Что конкретно должно быть приложено, в каком виде", "Комментарий"]
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    assert [r[0] for r in rows] == ["1.1", "6.8–6.12, 7.1.7"]  # только kind == "registry"
    assert rows[0][2] == "Столы круглые — 12 шт" and rows[0][3] == "Столы круглые — 12 шт. Расстановка по схеме."
    note = ws.cell(row=2, column=6)
    assert note.data_type == "s" and note.value.startswith("=")
    assert note.fill.fgColor.rgb.endswith("FFF8C5")
    assert ws.cell(row=3, column=6).fill.fgColor.rgb in ("00000000", None) or ws.cell(row=3, column=6).fill.fill_type is None
    assert ws.cell(row=1, column=1).font.b and ws.cell(row=2, column=4).alignment.wrap_text
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref == "A1:F3"
    assert not ws.merged_cells.ranges
    assert ws.page_setup.orientation == "landscape" and int(ws.page_setup.paperSize) == 9
    assert ws.sheet_properties.pageSetUpPr.fitToPage and ws.page_setup.fitToWidth == 1
    ws2 = wb["Требуют уточнения"]
    assert [c.value for c in ws2[1]] == ["Пункт ТЗ", "Раздел", "Наименование", "Описание",
                                          "Причина неопределённости"]
    assert [c.value for c in ws2[2]] == ["6.40", "6. Требования к услугам", "Брендирование",
                                          "Объём брендирования не указан", "+нет количества"]
    assert ws2.cell(row=2, column=5).data_type == "s"
    with zipfile.ZipFile(out) as z:
        assert "<f>" not in z.read("xl/worksheets/sheet1.xml").decode()

    # без kind — все пункты; без clarify — один лист
    plain = {"name": "П", "items": [{"id": "a", "num": "1", "title": "Стол", "text": "Стол"},
                                    {"id": "b", "num": ["2.1", "2.2"], "title": "Стулья", "text": "Стулья белые"}]}
    wb = openpyxl.load_workbook(export_registry_xlsx(plain, str(tmp_path / "r2.xlsx")))
    assert wb.sheetnames == ["Реестр подтверждений"]
    rows = [[c.value for c in r] for r in wb.active.iter_rows(min_row=2)]
    assert rows == [["1", None, "Стол", None, None, None], ["2.1, 2.2", None, "Стулья", "Стулья белые", None, None]]
