"""Разбор ТЗ в формате Word (.docx).

Особенность Word: автонумерация («1.», «1.2.», «а)») не хранится в тексте абзаца.
Её восстанавливаем сами по numbering.xml: у абзаца (или его стиля, например
«List Number» или заголовка, привязанного к многоуровневому списку) есть numId и ilvl,
по ним находим формат уровня (lvlText «%1.%2.», numFmt, start) и ведём счётчики
в порядке документа — так же, как это делает Word.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.table import Table, _Cell
from docx.text.hyperlink import Hyperlink
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from .common import (Para, TableBlock, build_items, clean_text, esc, src_cell, src_para, src_row,
                     title_from_filename)

_RU_LETTERS = "абвгдежзиклмнопрстуфхцчшщэюя"  # как russianLower в Word (без ё, й, ъ, ы, ь)
_EN_LETTERS = "abcdefghijklmnopqrstuvwxyz"


def _wval(el, tag: str) -> str | None:
    """Значение w:val дочернего элемента tag (или None)."""
    if el is None:
        return None
    c = el.find(qn(tag))
    if c is None:
        return None
    return c.get(qn("w:val"))


def _roman(n: int) -> str:
    vals = [(1000, "M"), (900, "CM"), (500, "D"), (400, "CD"), (100, "C"), (90, "XC"),
            (50, "L"), (40, "XL"), (10, "X"), (9, "IX"), (5, "V"), (4, "IV"), (1, "I")]
    out = []
    for v, s in vals:
        while n >= v:
            out.append(s)
            n -= v
    return "".join(out) or "0"


def _letters(n: int, alphabet: str) -> str:
    k = len(alphabet)
    n = max(1, n)
    return alphabet[(n - 1) % k] * ((n - 1) // k + 1)


def format_counter(n: int, fmt: str | None) -> str:
    """Число в формате уровня списка Word (decimal, lowerLetter, russianLower, upperRoman…)."""
    if fmt in ("bullet", "none"):
        return ""
    if fmt == "decimalZero":
        return f"{n:02d}"
    if fmt == "upperRoman":
        return _roman(n)
    if fmt == "lowerRoman":
        return _roman(n).lower()
    if fmt == "upperLetter":
        return _letters(n, _EN_LETTERS).upper()
    if fmt == "lowerLetter":
        return _letters(n, _EN_LETTERS)
    if fmt == "russianUpper":
        return _letters(n, _RU_LETTERS).upper()
    if fmt == "russianLower":
        return _letters(n, _RU_LETTERS)
    return str(n)


@dataclass
class _Lvl:
    start: int = 1
    fmt: str = "decimal"
    text: str = "%1."
    restart: int | None = None
    is_lgl: bool = False
    pstyle: str | None = None


def _parse_lvl(el, ilvl: int) -> _Lvl:
    start = _wval(el, "w:start")
    fmt = _wval(el, "w:numFmt") or "decimal"
    lt = el.find(qn("w:lvlText"))
    text = lt.get(qn("w:val")) if lt is not None else None
    if text is None:
        text = "•" if fmt == "bullet" else f"%{ilvl + 1}."
    rs = _wval(el, "w:lvlRestart")
    lgl = el.find(qn("w:isLgl"))
    return _Lvl(
        start=int(start) if start and start.lstrip("-").isdigit() else 1,
        fmt=fmt,
        text=text,
        restart=int(rs) if rs and rs.isdigit() else None,
        is_lgl=lgl is not None and lgl.get(qn("w:val")) not in ("0", "false"),
        pstyle=_wval(el, "w:pStyle"),
    )


class WordNumbering:
    """Восстановление автонумерации Word в порядке документа."""

    def __init__(self, document) -> None:
        self.abstracts: dict[str, dict] = {}
        self.nums: dict[str, tuple[str, dict]] = {}
        self.counters: dict[str, list] = {}
        self.seen: set[str] = set()
        self.styles = {}
        try:
            for s in document.styles:
                self.styles[s.style_id] = s
        except Exception:
            pass
        try:
            root = document.part.numbering_part.element
        except Exception:
            root = None
        if root is None:
            return
        for an in root.findall(qn("w:abstractNum")):
            aid = an.get(qn("w:abstractNumId"))
            levels = {}
            for lvl in an.findall(qn("w:lvl")):
                il = int(lvl.get(qn("w:ilvl"), "0") or 0)
                levels[il] = _parse_lvl(lvl, il)
            self.abstracts[aid] = {"levels": levels, "link": _wval(an, "w:numStyleLink")}
        for n in root.findall(qn("w:num")):
            nid = n.get(qn("w:numId"))
            aid = _wval(n, "w:abstractNumId")
            overrides = {}
            for ov in n.findall(qn("w:lvlOverride")):
                il = int(ov.get(qn("w:ilvl"), "0") or 0)
                so = _wval(ov, "w:startOverride")
                lvl_el = ov.find(qn("w:lvl"))
                overrides[il] = {
                    "start": int(so) if so and so.lstrip("-").isdigit() else None,
                    "lvl": _parse_lvl(lvl_el, il) if lvl_el is not None else None,
                }
            self.nums[nid] = (aid, overrides)

    # -- поиск определений ------------------------------------------------------
    def _resolve_abstract(self, aid: str | None, depth: int = 0) -> str | None:
        """Список, ссылающийся на стиль нумерации (numStyleLink), — к реальному определению."""
        a = self.abstracts.get(aid or "")
        if a is None or depth > 5:
            return aid
        link = a.get("link")
        if link and not a["levels"]:
            st = self.styles.get(link)
            if st is not None and st.element.pPr is not None:
                np_ = st.element.pPr.find(qn("w:numPr"))
                nid = _wval(np_, "w:numId") if np_ is not None else None
                if nid and nid in self.nums:
                    return self._resolve_abstract(self.nums[nid][0], depth + 1)
        return aid

    def _level_for_style(self, num_id: str, style_id: str) -> int:
        aid, _ = self.nums.get(num_id, (None, {}))
        aid = self._resolve_abstract(aid)
        for il, lv in self.abstracts.get(aid or "", {}).get("levels", {}).items():
            if lv.pstyle == style_id:
                return il
        return 0

    def numpr(self, p: Paragraph) -> tuple[str, int] | None:
        """numId и ilvl абзаца с учётом стиля (и цепочки базовых стилей)."""
        num_id = None
        ilvl = None
        pPr = p._p.pPr
        if pPr is not None:
            np_ = pPr.find(qn("w:numPr"))
            if np_ is not None:
                num_id = _wval(np_, "w:numId")
                il = _wval(np_, "w:ilvl")
                ilvl = int(il) if il and il.isdigit() else None
        if num_id is None:
            try:
                st = p.style
            except Exception:
                st = None
            seen = set()
            while st is not None and st.style_id not in seen:
                seen.add(st.style_id)
                spPr = st.element.pPr
                np_ = spPr.find(qn("w:numPr")) if spPr is not None else None
                if np_ is not None and _wval(np_, "w:numId") is not None:
                    num_id = _wval(np_, "w:numId")
                    if ilvl is None:
                        il = _wval(np_, "w:ilvl")
                        ilvl = int(il) if il and il.isdigit() else self._level_for_style(num_id, st.style_id)
                    break
                st = st.base_style
        if num_id is None or num_id == "0" or num_id not in self.nums:
            return None
        return num_id, max(0, min(8, ilvl or 0))

    # -- счёт ---------------------------------------------------------------------
    def number(self, num_id: str, ilvl: int) -> tuple[str, bool]:
        """Следующий номер для абзаца списка: (текст номера, это маркер-точка)."""
        aid, ov = self.nums[num_id]
        aid = self._resolve_abstract(aid) or f"num{num_id}"
        levels = self.abstracts.get(aid, {}).get("levels", {})

        def lv(i: int) -> _Lvl:
            o = ov.get(i)
            if o and o["lvl"] is not None:
                return o["lvl"]
            return levels.get(i) or _Lvl(1, "decimal", f"%{i + 1}.")

        def start_of(i: int) -> int:
            # startOverride задаёт начало уровня для этого списка (w:num): и для счёта,
            # и для показа «%1», если сам уровень ещё не встречался (4.1 при старте с ilvl=1)
            o = ov.get(i)
            if o and o["start"] is not None:
                return o["start"]
            return lv(i).start

        cnt = self.counters.setdefault(aid, [None] * 9)
        if num_id not in self.seen:
            self.seen.add(num_id)
            for i, o in ov.items():
                if o["start"] is not None or o["lvl"] is not None:
                    cnt[i] = None
        L = lv(ilvl)
        cnt[ilvl] = start_of(ilvl) if cnt[ilvl] is None else cnt[ilvl] + 1
        for d in range(ilvl + 1, 9):
            r = lv(d).restart
            if r is None or (r != 0 and ilvl <= r - 1):
                cnt[d] = None
        if L.fmt == "bullet":
            return (L.text or "•").strip() or "•", True

        def rep(m: re.Match) -> str:
            k = int(m.group(1)) - 1
            Lk = lv(k)
            v = cnt[k] if cnt[k] is not None else start_of(k)
            fmt = "decimal" if L.is_lgl else Lk.fmt
            return format_counter(v, fmt)

        return re.sub(r"%([1-9])", rep, L.text or "").strip(), False

    def for_paragraph(self, p: Paragraph) -> tuple[str, int, bool] | None:
        np_ = self.numpr(p)
        if np_ is None:
            return None
        num, bullet = self.number(*np_)
        return num, np_[1], bullet


# ---------------------------------------------------------------------------
# Свойства абзаца
# ---------------------------------------------------------------------------

def _style_chain(p: Paragraph):
    try:
        st = p.style
    except Exception:
        return
    seen = set()
    while st is not None and st.style_id not in seen:
        seen.add(st.style_id)
        yield st
        st = st.base_style


def _heading_rank(p: Paragraph) -> tuple[int, bool, int]:
    """(ранг заголовка, стиль Title/Subtitle, вложенность стиля списка «List Number 2»)."""
    pPr = p._p.pPr
    ol = _wval(pPr, "w:outlineLvl") if pPr is not None else None
    rank, is_title, sub = 0, False, 0
    for st in _style_chain(p):
        name = (st.name or "").strip()
        m = re.match(r"^(heading|заголовок)\s*(\d)$", name, re.I)
        if m and not rank:
            rank = int(m.group(2))
        if re.match(r"^(title|subtitle|название|заголовок|подзаголовок)$", name, re.I):
            is_title = True
        m = re.match(r"^list\s*(number|bullet|continue|paragraph)?\s*(\d)$", name, re.I)
        if m and not sub:
            sub = max(0, int(m.group(2)) - 1)
        if ol is None and st.element.pPr is not None:
            ol = _wval(st.element.pPr, "w:outlineLvl")
    if not rank and ol is not None and ol.isdigit() and int(ol) < 9:
        rank = int(ol) + 1
    return rank, is_title, sub


def _style_bold(p: Paragraph) -> bool:
    for st in _style_chain(p):
        try:
            b = st.font.bold
        except Exception:
            b = None
        if b is not None:
            return bool(b)
    return False


def _iter_runs(p: Paragraph):
    try:
        content = list(p.iter_inner_content())
    except Exception:
        content = list(p.runs)
    for item in content:
        if isinstance(item, Hyperlink):
            yield from item.runs
        elif isinstance(item, Run):
            yield item


def _para_bold(p: Paragraph) -> bool:
    base = _style_bold(p)
    any_text = False
    for r in _iter_runs(p):
        if not r.text.strip():
            continue
        any_text = True
        b = r.bold
        if b is None:
            b = base
        if not b:
            return False
    return any_text


def _alignment(p: Paragraph) -> str:
    al = p.alignment
    if al is None:
        for st in _style_chain(p):
            try:
                al = st.paragraph_format.alignment
            except Exception:
                al = None
            if al is not None:
                break
    if al == WD_ALIGN_PARAGRAPH.CENTER:
        return "center"
    if al == WD_ALIGN_PARAGRAPH.RIGHT:
        return "right"
    return ""


def _runs_html(p: Paragraph) -> str:
    """Текст абзаца в HTML с жирным/курсивом/подчёркиванием; всё экранировано."""
    segs: list[list] = []
    for r in _iter_runs(p):
        t = r.text
        if not t:
            continue
        key = (bool(r.bold), bool(r.italic), bool(r.underline))
        if segs and segs[-1][1] == key:
            segs[-1][0] += t
        else:
            segs.append([t, key])
    out = []
    for t, (b, i, u) in segs:
        h = esc(t.replace("\t", " ")).replace("\n", "<br>")
        if u:
            h = f"<u>{h}</u>"
        if i:
            h = f"<i>{h}</i>"
        if b:
            h = f"<b>{h}</b>"
        out.append(h)
    html = "".join(out).strip()
    return html or esc(clean_text(p.text)).replace("\n", "<br>")


# ---------------------------------------------------------------------------
# Обход документа
# ---------------------------------------------------------------------------

def _iter_blocks(parent_el, parent_obj):
    """Абзацы и таблицы тела документа по порядку (включая элементы управления содержимым)."""
    for child in parent_el.iterchildren():
        tag = child.tag
        if tag == qn("w:p"):
            yield Paragraph(child, parent_obj)
        elif tag == qn("w:tbl"):
            yield Table(child, parent_obj)
        elif tag == qn("w:sdt"):
            content = child.find(qn("w:sdtContent"))
            if content is not None:
                yield from _iter_blocks(content, parent_obj)
        elif tag == qn("w:customXml"):
            yield from _iter_blocks(child, parent_obj)


def _plain_line(text: str, auto: tuple | None) -> str:
    """Абзац для full_text: с номером автонумерации Word, маркированный — «- …»."""
    if not auto:
        return text
    num, _ilvl, bullet = auto
    if bullet:
        return f"- {text}".rstrip()
    return f"{num} {text}".strip() if num else text


def _cell_content(cell: _Cell, numbering: WordNumbering, ti: int, ri: int, ci: int,
                  pk: int) -> tuple[str, str, bool, list[Para], int, str]:
    """Содержимое ячейки: текст (с автонумерацией), HTML, «вся жирная», абзацы (Para), счётчик абзацев
    и текст для full_text (каждый абзац — с новой строки).

    pk — сквозной номер абзаца в строке таблицы; каждый непустой абзац получает
    data-src="t<ti>-r<ri>-p<pk>" (на него ссылается source.ref пунктов из ячейки).
    Ручные переносы строк (w:br) сохраняются как перевод строки.
    """
    texts, htmls, bolds, paras, plain = [], [], [], [], []
    for p in cell.paragraphs:
        k = pk
        pk += 1
        t = clean_text(p.text)
        auto = numbering.for_paragraph(p)
        mark = ("•" if auto[2] else auto[0]) if auto else ""
        # пустой абзац с автономером — обычная колонка «№» в Word: номер и есть содержимое
        if not t and not (auto and not auto[2] and auto[0]):
            continue
        h = _runs_html(p) if t else ""
        shown = t
        if mark:
            shown = f"{mark} {t}".strip()
            h = f'<span class="tz-num">{esc(mark)}</span> {h}'.strip()
        _rank, _is_title, sub = _heading_rank(p)
        bold = _para_bold(p) if t else False
        if t:
            paras.append(Para(text=t, src=src_cell(ti, ri, ci, k), bold=bold, auto=auto, sub_list=sub))
        texts.append(shown)
        plain.append(_plain_line(t, auto))
        htmls.append(f'<p class="tz-cp" data-src="t{ti}-r{ri}-p{k}">{h}</p>')
        if t:
            bolds.append(bold)
    for inner in cell.tables:  # вложенные таблицы — построчно текстом
        for row in inner.rows:
            vals = [clean_text(c.text) for c in row.cells]
            line = " | ".join(dict.fromkeys(v for v in vals if v))
            if line:
                texts.append(line)
                plain.append(line)
                htmls.append(f'<p class="tz-cp">{esc(line)}</p>')
                bolds.append(False)
    return "\n".join(texts), "".join(htmls), bool(bolds) and all(bolds), paras, pk, "\n".join(plain)


def _table(table: Table, ti: int, numbering: WordNumbering) -> tuple[TableBlock, str, str]:
    """Таблица → (TableBlock, HTML, текст для full_text: строка таблицы — ячейки через « | »)."""
    rows: list[list[str]] = []
    plain_rows: list[str] = []
    cell_paras: list[list[list[Para]]] = []
    srcs, merged, bold = [], [], []
    html = [f'<table class="tz-table" data-table="{ti}"><tbody>']
    for ri, tr in enumerate(table._tbl.tr_lst):
        cells: list[str] = []
        cparas: list[list[Para]] = []
        tds: list[str] = []
        spans_text: list[tuple[str, int]] = []
        row_bold: list[bool] = []
        row_plain: list[str] = []
        pk = 0
        before = getattr(tr, "grid_before", 0) or 0
        for _ in range(before):
            cells.append("")
            cparas.append([])
            tds.append("<td></td>")
        for tc in tr.tc_lst:
            span = max(1, tc.grid_span or 1)
            vmerge = tc.vMerge
            cell = _Cell(tc, table)
            ci = len(cells)
            if vmerge == "continue":
                text, chtml, cbold, paras, cplain = "", "", False, [], ""
            else:
                text, chtml, cbold, paras, pk, cplain = _cell_content(cell, numbering, ti, ri, ci, pk)
            row_plain.append(cplain)
            cells.append(text)
            cells.extend([""] * (span - 1))
            cparas.append(paras)
            cparas.extend([[] for _ in range(span - 1)])
            spans_text.append((text, span))
            if text:
                row_bold.append(cbold)
            colspan = f' colspan="{span}"' if span > 1 else ""
            tds.append(f"<td{colspan}>{chtml}</td>")
        rows.append(cells)
        cell_paras.append(cparas)
        srcs.append(src_row(ti, ri))
        nonempty = [(t, s) for t, s in spans_text if t]
        total = len(cells)
        merged.append(len(nonempty) == 1 and total > 1 and nonempty[0][1] >= max(2, total - before))
        bold.append(bool(row_bold) and all(row_bold))
        html.append(f'<tr data-src="t{ti}-r{ri}">{"".join(tds)}</tr>')
        if any(row_plain):
            plain_rows.append(" | ".join(row_plain))
    html.append("</tbody></table>")
    return (TableBlock(rows=rows, srcs=srcs, merged=merged, bold=bold, cell_paras=cell_paras), "".join(html),
            "\n".join(plain_rows))


def paragraph_refs(document) -> list[tuple[str, object]]:
    """Абзацы документа тем же обходом, что parse_docx: [(ref, элемент w:p)] в порядке документа.

    ref совпадает с source.ref пунктов («p12», «t0-r6-p16»); для строки таблицы («t0-r6») — её последний
    абзац. Нужен выгрузке «ТЗ с фото», чтобы вставить отметки и фото в исходный файл под пунктами.
    """
    out: list[tuple[str, object]] = []
    p_index = t_index = 0
    for block in _iter_blocks(document.element.body, document._body):
        if isinstance(block, Paragraph):
            out.append((f"p{p_index}", block._p))
            p_index += 1
            continue
        ti = t_index
        t_index += 1
        for ri, tr in enumerate(block._tbl.tr_lst):
            pk, last = 0, None
            for tc in tr.tc_lst:
                if tc.vMerge == "continue":
                    continue
                for p in _Cell(tc, block).paragraphs:
                    out.append((f"t{ti}-r{ri}-p{pk}", p._p))
                    last = p._p
                    pk += 1
            if last is not None:
                out.append((f"t{ti}-r{ri}", last))
    return out


def parse_docx(data: bytes, filename: str) -> dict:
    """Разбор .docx → словарь parse_tz (см. app/tz/__init__.py)."""
    try:
        document = Document(io.BytesIO(data))
    except Exception as e:  # noqa: BLE001 — любые ошибки чтения пакета
        raise ValueError("Не удалось открыть файл Word: он повреждён или это не .docx") from e
    numbering = WordNumbering(document)
    stream: list = []
    html: list[str] = ['<div class="tz-doc tz-docx">']
    p_index = 0
    t_index = 0
    full: list[str] = []   # весь текст документа по порядку — для full_text
    for block in _iter_blocks(document.element.body, document._body):
        if isinstance(block, Paragraph):
            idx = p_index
            p_index += 1
            text = clean_text(block.text)
            auto = numbering.for_paragraph(block)
            if not text:
                continue
            full.append(_plain_line(text, auto))
            rank, is_title, sub = _heading_rank(block)
            align = _alignment(block)
            para = Para(text=text, src=src_para(idx), bold=_para_bold(block), heading=rank,
                        is_title=is_title, centered=(align == "center"), auto=auto, sub_list=sub)
            stream.append(para)
            inner = _runs_html(block)
            style = ""
            if auto:
                num, ilvl, bullet = auto
                mark = "•" if bullet else num
                if mark:
                    inner = f'<span class="tz-num">{esc(mark)}</span> {inner}'
                indent = ilvl + sub
                if indent:
                    style = f' style="margin-left:{1.5 * indent:.1f}em"'
            if align:
                style = style[:-1] + f';text-align:{align}"' if style else f' style="text-align:{align}"'
            if is_title or rank:
                tag = "h1" if is_title and not rank else f"h{min(6, rank + 1)}"
                html.append(f'<{tag} class="tz-h" data-src="p{idx}"{style}>{inner}</{tag}>')
            else:
                cls = "tz-li" if auto else "tz-p"
                html.append(f'<p class="{cls}" data-src="p{idx}"{style}>{inner}</p>')
        else:
            tb, thtml, tplain = _table(block, t_index, numbering)
            t_index += 1
            stream.append(tb)
            html.append(thtml)
            if tplain:
                full.extend(["", tplain, ""])
    html.append("</div>")
    title, items = build_items(stream, title_from_filename(filename))
    warnings = []
    if not items:
        warnings.append("В документе не найдено пунктов: добавьте их вручную.")
    full_text = re.sub(r"\n{3,}", "\n\n", "\n".join(full)).strip()
    return {"title": title, "kind": "docx", "items": items, "html": "\n".join(html),
            "pages": [], "warnings": warnings, "full_text": full_text}
