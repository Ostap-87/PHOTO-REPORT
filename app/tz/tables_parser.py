"""Разбор ТЗ из таблиц: Excel (.xlsx, .xls) и CSV.

Ищем лист с перечнем работ: строку заголовков определяем по синонимам
(«№ п/п», «Наименование работ», «Кол-во», «Ед. изм.» …), строки выше — шапка
и название, строки ниже — пункты. Если перечня нет, каждая строка листа
разбирается как абзац текста (маркеры «1.», «1.2» и т. п.).
"""

from __future__ import annotations

import csv
import datetime as _dt
import io

from .common import (Para, TableBlock, build_items, decode_text, find_table_map, one_line, src_row,
                     table_html, title_from_filename)

MAX_ROWS = 5000
MAX_COLS = 40


def _cell_str(v) -> str:
    """Значение ячейки строкой (даты — ДД.ММ.ГГГГ, целые числа без «.0»)."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "да" if v else "нет"
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else repr(v)
    if isinstance(v, _dt.datetime):
        return v.strftime("%d.%m.%Y %H:%M") if (v.hour or v.minute) else v.strftime("%d.%m.%Y")
    if isinstance(v, _dt.date):
        return v.strftime("%d.%m.%Y")
    if isinstance(v, _dt.time):
        return v.strftime("%H:%M")
    return str(v).strip()


def _trim(rows: list[list[str]]) -> list[list[str]]:
    """Убирает пустые строки в конце и пустые колонки справа."""
    while rows and not any(c.strip() for c in rows[-1]):
        rows.pop()
    width = 0
    for r in rows:
        for i in range(len(r) - 1, -1, -1):
            if r[i].strip():
                width = max(width, i + 1)
                break
    return [(r + [""] * width)[:width] for r in rows]


class _Sheet:
    def __init__(self, name: str, rows: list[list[str]], spans: dict | None = None):
        self.name = name
        self.rows = _trim(rows)
        self.spans = {k: v for k, v in (spans or {}).items()
                      if k[0] < len(self.rows) and (not self.rows or k[1] < len(self.rows[0]))}


def _merged_rows(sheet: _Sheet) -> list[bool]:
    """Строка — одна объединённая по ширине ячейка (строка-раздел)."""
    out = []
    for r, row in enumerate(sheet.rows):
        vals = [c for c in row if c.strip()]
        wide = any(rr == r and cs >= 2 for (rr, _c), (_rs, cs) in sheet.spans.items())
        out.append(len(vals) == 1 and wide)
    return out


def _build(sheets: list[_Sheet], kind: str, filename: str) -> dict:
    html = ['<div class="tz-doc tz-sheet">']
    stream: list = []
    work = []
    for k, sh in enumerate(sheets):
        html.append(table_html(sh.rows, k, sh.spans, caption=sh.name if len(sheets) > 1 else ""))
        tmap = find_table_map(sh.rows)
        if tmap is not None:
            work.append((k, sh, tmap))

    if work:
        for n, (k, sh, tmap) in enumerate(work):
            start = max(0, tmap.header_row)
            if n == 0:
                for r in range(start):
                    text = " ".join(c for c in sh.rows[r] if c.strip())
                    if text:
                        stream.append(Para(text=text, src=src_row(k, r)))
            if len(work) > 1:
                stream.append(Para(text=sh.name, src=src_row(k, start), heading=1))
            merged = _merged_rows(sh)
            stream.append(TableBlock(rows=sh.rows[start:], srcs=[src_row(k, r) for r in range(start, len(sh.rows))],
                                     merged=merged[start:], name=sh.name))
    else:
        for k, sh in enumerate(sheets):
            for r, row in enumerate(sh.rows):
                vals = [c for c in row if c.strip()]
                if vals:
                    stream.append(Para(text=" — ".join(vals), src=src_row(k, r)))
    html.append("</div>")
    title, items = build_items(stream, title_from_filename(filename))
    warnings = []
    if not work:
        warnings.append("Не найдена таблица с колонкой «Наименование»: строки разобраны как текст.")
    if not items:
        warnings.append("В таблице не найдено пунктов: добавьте их вручную.")
    return {"title": title, "kind": kind, "items": items, "html": "\n".join(html), "pages": [],
            "warnings": warnings}


def parse_xlsx(data: bytes, filename: str) -> dict:
    """Разбор .xlsx/.xlsm (openpyxl)."""
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except Exception as e:  # noqa: BLE001
        raise ValueError("Не удалось открыть файл Excel: он повреждён или это не .xlsx") from e
    sheets = []
    for ws in wb.worksheets:
        if getattr(ws, "sheet_state", "visible") != "visible":
            continue
        max_r = min(ws.max_row or 0, MAX_ROWS)
        max_c = min(ws.max_column or 0, MAX_COLS)
        if not max_r or not max_c:
            continue
        rows = [[_cell_str(c.value) for c in row]
                for row in ws.iter_rows(min_row=1, max_row=max_r, max_col=max_c)]
        spans = {}
        for mr in ws.merged_cells.ranges:
            if mr.min_row <= max_r and mr.min_col <= max_c:
                spans[(mr.min_row - 1, mr.min_col - 1)] = (
                    min(mr.max_row, max_r) - mr.min_row + 1, min(mr.max_col, max_c) - mr.min_col + 1)
        sh = _Sheet(ws.title, rows, spans)
        if sh.rows:
            sheets.append(sh)
    wb.close()
    if not sheets:
        raise ValueError("В файле Excel нет заполненных листов")
    return _build(sheets, "xlsx", filename)


def parse_xls(data: bytes, filename: str) -> dict:
    """Разбор старого .xls (нужен пакет xlrd)."""
    try:
        import xlrd
    except ImportError as e:
        raise ValueError("Для старого формата .xls нужен пакет xlrd. Либо сохраните файл как .xlsx.") from e
    try:
        try:
            book = xlrd.open_workbook(file_contents=data, formatting_info=True)
        except Exception:  # noqa: BLE001 — formatting_info поддерживается не для всех файлов
            book = xlrd.open_workbook(file_contents=data)
    except Exception as e:  # noqa: BLE001
        raise ValueError("Не удалось открыть файл .xls") from e
    sheets = []
    for sh in book.sheets():
        if getattr(sh, "visibility", 0):
            continue
        rows = []
        for r in range(min(sh.nrows, MAX_ROWS)):
            row = []
            for c in range(min(sh.ncols, MAX_COLS)):
                cell = sh.cell(r, c)
                v = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        v = xlrd.xldate_as_datetime(v, book.datemode)
                    except Exception:  # noqa: BLE001
                        pass
                row.append(_cell_str(v))
            rows.append(row)
        spans = {}
        for (r0, r1, c0, c1) in getattr(sh, "merged_cells", []) or []:
            spans[(r0, c0)] = (r1 - r0, c1 - c0)
        s = _Sheet(sh.name, rows, spans)
        if s.rows:
            sheets.append(s)
    if not sheets:
        raise ValueError("В файле .xls нет заполненных листов")
    return _build(sheets, "xlsx", filename)


def parse_csv(data: bytes, filename: str) -> dict:
    """Разбор CSV/TSV: кодировка и разделитель определяются автоматически."""
    text = decode_text(data)
    sample = text[:20000]
    delim = ";"
    try:
        delim = csv.Sniffer().sniff(sample, delimiters=";,\t|").delimiter
    except csv.Error:
        counts = {d: sample.count(d) for d in (";", "\t", ",", "|")}
        delim = max(counts, key=counts.get) if any(counts.values()) else ";"
    rows = [[one_line(c) for c in row[:MAX_COLS]] for row in csv.reader(io.StringIO(text), delimiter=delim)]
    rows = rows[:MAX_ROWS]
    sh = _Sheet("CSV", rows)
    if not sh.rows:
        raise ValueError("Файл CSV пустой")
    return _build([sh], "csv", filename)
