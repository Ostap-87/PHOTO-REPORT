"""«ТЗ с фото»: исходный файл ТЗ (Word), в который под пунктами вставлены отметка исполнения, фото,
файлы и ссылки — документ остаётся в оформлении заказчика, ИИ не нужен.

Место вставки находится по source.ref пункта (тот же обход документа, что при разборе ТЗ,
app.tz.docx_parser.paragraph_refs): после текста пункта, перед следующим пунктом в том же месте
документа (в тексте — перед следующим абзацем-пунктом или таблицей с пунктами, в таблице — в той же
ячейке). Строки реестра берут место своего пункта ТЗ. Пункты без места в ТЗ (добавленные вручную) и,
по флажку, «Неразобранное» — в конце документа.
"""

from __future__ import annotations

import difflib
import io
import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from app.tz.docx_parser import paragraph_refs
from app.tz.perfect import perfect

from .common import (STATUS_COLORS, STATUS_TINTS, ImageCache, PhotoIndex, file_label, fmt_date, item_extra,
                     item_label, norm_options, photo_caption, select_items, status_key, status_label)
from .docx_report import _cell_margins, _hyperlink, _para_box, _row_flags, _run, _widths

TWIP_CM = 567.0


def _rewrite_paragraph(p_el) -> bool:
    """Текст абзаца — «в совершенном виде» с сохранением оформления: правки по символам внутри w:t."""
    ts = list(p_el.iter(qn("w:t")))
    if not ts:
        return False
    segs = [t.text or "" for t in ts]
    full = "".join(segs)
    new = perfect(full)
    if new == full:
        return False
    ops = difflib.SequenceMatcher(None, full, new, autojunk=False).get_opcodes()
    for tag, i1, i2, j1, j2 in reversed(ops):
        if tag == "equal":
            continue
        # сегмент и смещение начала правки
        pos, k = 0, 0
        while k < len(segs) - 1 and pos + len(segs[k]) < i1:
            pos += len(segs[k])
            k += 1
        off = i1 - pos
        rest = i2 - i1
        seg = segs[k]
        segs[k] = seg[:off] + new[j1:j2] + seg[off + min(rest, len(seg) - off):]
        rest -= min(rest, len(seg) - off)
        n = k + 1
        while rest > 0 and n < len(segs):              # удаление, заходящее в следующие сегменты
            cut = min(rest, len(segs[n]))
            segs[n] = segs[n][cut:]
            rest -= cut
            n += 1
    for t, v in zip(ts, segs):
        if t.text != v:
            t.text = v
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    return True


def rewrite_perfect(doc) -> int:
    """Весь текст ТЗ (абзацы и ячейки таблиц) — в совершенном виде. -> сколько абзацев изменено."""
    return sum(_rewrite_paragraph(p) for p in list(doc.element.body.iter(qn("w:p"))))


class NotWordTZ(ValueError):
    """ТЗ проекта не в формате Word — «ТЗ с фото» не собрать (есть «Отчёт Word»)."""


def _usable_cm(doc) -> float:
    sec = doc.sections[-1]
    try:
        return max(8.0, (sec.page_width - sec.left_margin - sec.right_margin) / 360000)
    except TypeError:
        return 16.5


def _cell_cm(tc, doc) -> float:
    """Ширина ячейки (tcW в twips) или доля ширины страницы."""
    tcpr = tc.tcPr
    w = tcpr.find(qn("w:tcW")) if tcpr is not None else None
    try:
        if w is not None and w.get(qn("w:type")) in (None, "dxa") and int(w.get(qn("w:w"))) > 0:
            return int(w.get(qn("w:w"))) / TWIP_CM
    except (TypeError, ValueError):
        pass
    tr = tc.getparent()
    n = max(1, len(tr.findall(qn("w:tc"))))
    return _usable_cm(doc) / n


def _ancestor(el, tag: str):
    while el is not None:
        if el.tag == qn(tag):
            return el
        el = el.getparent()
    return None


def _target(refs: list, a: int, anchors: list[int]):
    """Куда вставлять для пункта с якорем refs[a]: ("before", элемент) или ("end", контейнер)."""
    el = refs[a][1]
    box = el.getparent()                                     # w:body / w:tc / w:sdtContent
    for n in anchors:
        if n <= a or refs[n][1] is el:
            continue
        cur = refs[n][1]
        while cur is not None and cur.getparent() is not box:  # подняться до соседа в том же контейнере
            cur = cur.getparent()
        if cur is not None:
            return ("before", cur)
        if box.tag == qn("w:tc"):                           # следующий пункт в другой ячейке — до конца своей
            return ("end", box)
    return ("end", box)


class _Builder:
    """Создаёт абзацы и таблицы в документе и переносит их на место вставки."""

    def __init__(self, doc, project: dict, project_dir: str, opts: dict):
        self.doc = doc
        self.opts = opts
        self.idx = PhotoIndex(project, opts["skip_duplicates"])
        self.cache = ImageCache(project_dir, opts["max_px"], opts["jpeg_quality"])
        self.counter = 0

    def has_content(self, it: dict) -> bool:
        ex = item_extra(it)
        return bool(status_key(it) != "todo" or self.idx.for_item(it) or self.idx.files_for(it) or ex["links"]
                    or (it.get("comment") or "").strip())

    def _para(self, text: str = "", size: float = 9, color: str | None = None, bold: bool = False):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        if text:
            _run(p, text, size=size, bold=bold, color=color)
        return p

    def blocks(self, it: dict, width_cm: float, with_label: bool = False) -> list:
        """Элементы для пункта: отметка исполнения, фото, файлы, ссылки (созданы в конце документа)."""
        out = []
        key = status_key(it)
        ex = item_extra(it)
        p = self._para()
        if with_label:
            _run(p, item_label(it) + "\n", size=9.5, bold=True, color="202124")
        _run(p, "Исполнение: ", size=9, bold=True, color="3C4043")
        _run(p, status_label(key), size=9, bold=True, color=STATUS_COLORS[key])
        if it.get("done_at") and key != "todo":
            _run(p, " · " + fmt_date(it["done_at"]), size=9, color="5F6368")
        if (it.get("comment") or "").strip():
            _run(p, " — " + it["comment"].strip(), size=9, color="3C4043")
        if ex["required"] and it.get("kind") == "registry":
            _run(p, "\nПодтверждение: " + ex["required"], size=8.5, color="5F6368")
        _para_box(p, STATUS_TINTS[key], STATUS_COLORS[key])
        out.append(p._p)
        photos = self.idx.for_item(it)
        if photos:
            t = self._grid(photos, width_cm)
            if t is not None:
                out.append(t._tbl)
        files = self.idx.files_for(it)
        if files:
            p = self._para("Файлы: " + "; ".join(file_label(f) for f in files), color="3C4043")
            out.append(p._p)
        for url in ex["links"]:
            p = self._para()
            _run(p, "Ссылка: ", size=9, color="3C4043")
            try:
                _hyperlink(p, url, url, size=9)
            except Exception:  # noqa: BLE001 — не ссылка http(s): просто текст
                _run(p, url, size=9, color="1A73E8")
            out.append(p._p)
        return out

    def _grid(self, photos: list[dict], width_cm: float):
        per_row = self.opts["photos_per_row"] if width_cm >= 11 else 1
        prepared = []
        for ph in photos:
            img = self.cache.get(ph)
            if img:
                self.counter += 1
                prepared.append((ph, img, self.counter))
        if not prepared:
            return None
        cell_w = width_cm / per_row
        img_h = {1: 10.0, 2: 7.0, 3: 5.0, 4: 4.0}[per_row]
        t = self.doc.add_table(rows=(len(prepared) + per_row - 1) // per_row, cols=per_row)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _cell_margins(t, 30, 60, 40, 40)
        for i, (ph, (data, w, h), n) in enumerate(prepared):
            cell = t.rows[i // per_row].cells[i % per_row]
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(io.BytesIO(data), width=Cm(max(2.0, min(cell_w - 0.4, img_h * w / max(1, h)))))
            cap = cell.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _run(cap, photo_caption(n, ph), size=7.5, color="5F6368")
        for row in t.rows:
            _row_flags(row)
        _widths(t, [cell_w] * per_row)
        return t


def _place(doc, where, elements: list) -> None:
    mode, ref = where
    for el in elements:
        el.getparent().remove(el)
    if mode == "before":
        for el in elements:
            ref.addprevious(el)
        return
    box = ref
    sect = box.find(qn("w:sectPr")) if box.tag == qn("w:body") else None
    for el in elements:
        if sect is not None:
            sect.addprevious(el)
        else:
            box.append(el)
    if box.tag == qn("w:tc") and box[-1].tag != qn("w:p"):    # ячейка Word обязана кончаться абзацем
        box.append(doc.add_paragraph()._p)


def export_tz_docx(project: dict, project_dir: str, out_path: str, options: dict | None = None) -> str:
    """Собрать «ТЗ с фото» из исходного .docx проекта; возвращает out_path."""
    opts = norm_options(options)
    tz = project.get("tz") or {}
    src = os.path.join(project_dir, tz.get("file") or "")
    if tz.get("kind") != "docx" or not tz.get("file") or not os.path.isfile(src):
        raise NotWordTZ("«ТЗ с фото» собирается из ТЗ в формате Word (.docx). Для ТЗ из PDF, Excel или текста "
                        "выгрузите «Отчёт Word».")
    doc = Document(src)
    refs = paragraph_refs(doc)
    if opts.get("perfect", True):                            # «Исполнитель обеспечил…» вместо «обеспечивает»
        rewrite_perfect(doc)
    pos: dict[str, int] = {}
    for i, (ref, _el) in enumerate(refs):
        pos.setdefault(ref, i)
    tz_items = project.get("tz_items") or project.get("items") or []
    anchors = sorted({pos[r] for it in tz_items if (r := (it.get("source") or {}).get("ref")) in pos})
    b = _Builder(doc, project, project_dir, opts)

    placed: dict[int, list[dict]] = {}
    loose: list[dict] = []
    for it in select_items(project, opts):
        if not b.has_content(it):
            continue
        r = (it.get("source") or {}).get("ref")
        if r in pos:
            placed.setdefault(pos[r], []).append(it)
        else:
            loose.append(it)

    for a in sorted(placed):                                 # по порядку документа — сквозная нумерация фото
        where = _target(refs, a, anchors)
        box = where[1] if where[0] == "end" else where[1].getparent()
        tc = box if box.tag == qn("w:tc") else _ancestor(box, "w:tc")
        width = _cell_cm(tc, doc) - 0.4 if tc is not None else _usable_cm(doc)
        els = []
        for it in placed[a]:
            els += b.blocks(it, width, with_label=len(placed[a]) > 1 or it.get("kind") == "registry")
        _place(doc, where, els)

    loose_photos = b.idx.unsorted() if opts["include_unsorted"] else []
    if loose or loose_photos:                               # add_* дописывает в конец документа — там и остаются
        p = b._para()
        p.add_run().add_break(WD_BREAK.PAGE)
        _run(p, "Дополнительно к ТЗ", size=13, bold=True, color="202124")
        for it in loose:
            b.blocks(it, _usable_cm(doc), with_label=True)
        if loose_photos:
            b._para(f"Фото без пункта ТЗ: {len(loose_photos)}", size=10, bold=True, color="3C4043")
            b._grid(loose_photos, _usable_cm(doc))

    cp = doc.core_properties
    cp.title = f"ТЗ с фото — {project.get('name') or ''}".strip(" —")
    cp.comments = "Собрано программой ФотоОтчёт: отметки исполнения и фото под пунктами ТЗ"
    os.makedirs(os.path.dirname(os.path.abspath(out_path)) or ".", exist_ok=True)
    doc.save(out_path)
    return out_path
