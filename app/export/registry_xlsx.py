"""Реестр подтверждений (Excel): что и в каком виде приложить по каждому пункту ТЗ.

    export_registry_xlsx(project, out_path) -> str

Лист «Реестр подтверждений»: Пункт ТЗ | Раздел | Наименование | Описание |
Что конкретно должно быть приложено, в каком виде | Комментарий — из пунктов с kind == "registry"
(если таких нет — из всех пунктов). Лист «Требуют уточнения» (project["clarify"]) — только если
он не пуст. Без объединённых ячеек; формулоподобный текст сохраняется как текст.
"""

from __future__ import annotations

import os

from .common import item_extra, same_text
from .xlsx_report import _put

REGISTRY_COLUMNS = [
    ("num", "Пункт ТЗ", 13),
    ("section", "Раздел", 24),
    ("title", "Наименование", 38),
    ("text", "Описание", 60),
    ("required", "Что конкретно должно быть приложено, в каком виде", 52),
    ("note", "Комментарий", 40),
]
CLARIFY_COLUMNS = [
    ("num", "Пункт ТЗ", 13),
    ("section", "Раздел", 24),
    ("name", "Наименование", 38),
    ("description", "Описание", 60),
    ("reason", "Причина неопределённости", 52),
]


def _num_text(v) -> str:
    """Номер пункта: строка («6.34.3, 7.1.5») или список номеров."""
    if isinstance(v, (list, tuple)):
        return ", ".join(str(x).strip() for x in v if str(x).strip())
    return str(v or "").strip()


def _registry_rows(project: dict) -> list[dict]:
    items = [it for it in (project.get("items") or []) if isinstance(it, dict)]
    reg = [it for it in items if it.get("kind") == "registry"] or items
    rows = []
    for it in reg:
        ex = item_extra(it)
        title = (it.get("title") or "").strip()
        text = (it.get("text") or "").strip()
        rows.append({
            "num": _num_text(it.get("num")),
            "section": ex["section"],
            "title": title or text.split("\n")[0],
            "text": "" if (not text or same_text(text, title)) else text,
            "required": ex["required"],
            "note": ex["note"],
        })
    return rows


def _clarify_rows(project: dict) -> list[dict]:
    rows = []
    for c in project.get("clarify") or []:
        if not isinstance(c, dict):
            continue
        rows.append({k: (_num_text(c.get(k)) if k == "num" else str(c.get(k) or "").strip())
                     for k, _h, _w in CLARIFY_COLUMNS})
    return rows


def _fill_sheet(ws, columns: list[tuple], rows: list[dict], highlight: str | None = None) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.properties import PageSetupProperties

    thin = Side(style="thin", color="BDC1C6")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="E8F0FE")
    note_fill = PatternFill("solid", fgColor="FFF8C5")
    wrap = Alignment(vertical="top", wrap_text=True)
    for c, (_key, title, width) in enumerate(columns, start=1):
        cell = _put(ws, 1, c, title)
        cell.font = Font(bold=True, color="202124")
        cell.fill = head_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.row_dimensions[1].height = 32
    for r, row in enumerate(rows, start=2):
        for c, (key, _title, _w) in enumerate(columns, start=1):
            value = row.get(key) or ""
            cell = _put(ws, r, c, value)
            cell.border = border
            cell.alignment = wrap
            if key == highlight and value:
                cell.fill = note_fill
        ws.cell(row=r, column=1).font = Font(bold=True)
    last_col = get_column_letter(len(columns))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{last_col}{max(1, len(rows)) + 1}"
    # печать: A4 альбомная, по ширине страницы, заголовок на каждой странице
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:1"
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5


def export_registry_xlsx(project: dict, out_path: str) -> str:
    """Сохраняет реестр подтверждений (.xlsx); возвращает out_path."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Реестр подтверждений"
    _fill_sheet(ws, REGISTRY_COLUMNS, _registry_rows(project), highlight="note")
    clarify = _clarify_rows(project)
    if clarify:
        _fill_sheet(wb.create_sheet("Требуют уточнения"), CLARIFY_COLUMNS, clarify)
    wb.properties.title = f"Реестр подтверждений — {project.get('name') or ''}".strip(" —")
    wb.properties.creator = "ФотоОтчёт"
    os.makedirs(os.path.dirname(os.path.abspath(out_path)) or ".", exist_ok=True)
    wb.save(out_path)
    return out_path
