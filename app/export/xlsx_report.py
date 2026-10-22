"""Отчёт Excel: лист «Исполнение» (пункты ТЗ) и лист «Фото» (реестр снимков)."""

from __future__ import annotations

import os

from .common import (STATUS_COLORS, STATUS_TINTS, PhotoIndex, item_extra, parse_dt, photo_datetime, quality_text,
                     status_key, status_label)

_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


def _put(ws, row: int, col: int, value):
    """Записывает значение; строки, похожие на формулу, сохраняются как текст."""
    cell = ws.cell(row=row, column=col)
    if isinstance(value, str):
        value = value.replace("\x00", "")
        cell.value = value
        if value.startswith(_FORMULA_START):
            cell.data_type = "s"
            cell.quotePrefix = True
    else:
        cell.value = value
    return cell


def _yes_no(v) -> str:
    return "Да" if v else "Нет"


def export_xlsx(project: dict, project_dir: str, out_path: str) -> str:
    """Сохраняет .xlsx с листами «Исполнение» и «Фото»; возвращает out_path."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    items = project.get("items") or []
    idx = PhotoIndex(project)
    item_by_id = {it.get("id"): it for it in items}

    wb = Workbook()
    ws = wb.active
    ws.title = "Исполнение"
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="37474F")
    thin = Side(style="thin", color="D0D4D8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap_top = Alignment(vertical="top", wrap_text=True)

    extras = [item_extra(it) for it in items]
    has = {k: any(e[k] for e in extras) for k in ("section", "required", "note", "links")}
    has["files"] = bool(idx.files)
    # (ключ, заголовок, ширина); необязательные колонки — только если в проекте есть такие данные
    cols = [("num", "№", 8)]
    if has["section"]:
        cols.append(("section", "Раздел", 24))
    cols.append(("title", "Пункт ТЗ", 60))
    if has["required"]:
        cols.append(("required", "Что приложить", 45))
    cols += [("qty", "Кол-во", 10), ("unit", "Ед.", 8), ("status", "Статус", 15), ("date", "Дата", 12),
             ("comment", "Комментарий", 40)]
    if has["note"]:
        cols.append(("note", "Комментарий анализа", 40))
    if has["links"]:
        cols.append(("links", "Ссылки", 40))
    cols.append(("photos", "Фото, шт.", 9))
    if has["files"]:
        cols.append(("files", "Файлов", 9))
    cols.append(("photo_names", "Файлы фото", 50))
    col = {key: i for i, (key, _h, _w) in enumerate(cols, start=1)}
    for c, (_key, h, w) in enumerate(cols, start=1):
        cell = _put(ws, 1, c, h)
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
        ws.column_dimensions[get_column_letter(c)].width = w

    for r, (it, ex) in enumerate(zip(items, extras), start=2):
        st = status_key(it)
        photos = idx.for_item(it)
        level = int(it.get("level") or 1)
        done = parse_dt(it.get("done_at"))
        values = {
            "num": it.get("num") or "", "section": ex["section"], "title": it.get("title") or it.get("text") or "",
            "required": ex["required"], "qty": it.get("qty"), "unit": it.get("unit") or "",
            "status": status_label(st), "date": done if done else (it.get("done_at") or ""),
            "comment": it.get("comment") or "", "note": ex["note"], "links": "\n".join(ex["links"]),
            "photos": len(photos), "files": len(idx.files_for(it)),
            "photo_names": "; ".join(p.get("name") or p.get("id") for p in photos),
        }
        for key, c in col.items():
            cell = _put(ws, r, c, values[key])
            cell.border = border
            cell.alignment = wrap_top
        title_cell = ws.cell(row=r, column=col["title"])
        title_cell.alignment = Alignment(vertical="top", wrap_text=True, indent=max(0, level - 1) * 2)
        if level == 1:
            title_cell.font = Font(bold=True)
            ws.cell(row=r, column=col["num"]).font = Font(bold=True)
        st_cell = ws.cell(row=r, column=col["status"])
        st_cell.fill = PatternFill("solid", fgColor=STATUS_TINTS[st])
        st_cell.font = Font(bold=True, color=STATUS_COLORS[st])
        st_cell.alignment = Alignment(horizontal="center", vertical="top")
        if done:
            ws.cell(row=r, column=col["date"]).number_format = "DD.MM.YYYY"
        for key in ("photos", "files"):
            if key in col:
                ws.cell(row=r, column=col[key]).alignment = Alignment(horizontal="center", vertical="top")
        if "required" in col and ex["required"]:
            ws.cell(row=r, column=col["required"]).fill = PatternFill("solid", fgColor="FEF7E0")
    ws.freeze_panes = "A2"
    if items:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(items) + 1}"

    # Лист «Фото»
    ws2 = wb.create_sheet("Фото")
    headers2 = ["Файл", "Пункт ТЗ", "Дата съёмки", "GPS", "Камера", "Качество", "Улучшено", "Лица скрыты",
                "Уверенность автосортировки", "Подтверждено"]
    widths2 = [28, 50, 17, 22, 22, 22, 11, 12, 14, 13]
    for c, (h, w) in enumerate(zip(headers2, widths2), start=1):
        cell = _put(ws2, 1, c, h)
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
        ws2.column_dimensions[get_column_letter(c)].width = w

    ordered, seen = [], set()
    for it in items:
        for p in idx.for_item(it):
            if p["id"] not in seen:
                ordered.append(p)
                seen.add(p["id"])
    ordered += [p for p in idx.photos if p["id"] not in seen]

    for r, p in enumerate(ordered, start=2):
        it = item_by_id.get(p.get("item_id"))
        item_text = f"{it.get('num') or ''} {it.get('title') or ''}".strip() if it else "Неразобранное"
        exif = p.get("exif") or {}
        gps = exif.get("gps")
        gps_text = ""
        if isinstance(gps, (list, tuple)) and len(gps) >= 2 and gps[0] is not None and gps[1] is not None:
            try:
                gps_text = f"{float(gps[0]):.6f}, {float(gps[1]):.6f}"
            except (TypeError, ValueError):
                gps_text = str(gps)
        dt = photo_datetime(p)
        conf = p.get("confidence")
        values = [p.get("name") or p.get("id"), item_text, dt if dt else (exif.get("datetime") or ""), gps_text,
                  exif.get("camera") or "", quality_text(p), _yes_no(p.get("enhanced")),
                  _yes_no(p.get("faces_blurred")),
                  float(conf) if isinstance(conf, (int, float)) else (conf or ""), _yes_no(p.get("confirmed"))]
        for c, v in enumerate(values, start=1):
            cell = _put(ws2, r, c, v)
            cell.border = border
            cell.alignment = wrap_top
        if dt:
            ws2.cell(row=r, column=3).number_format = "DD.MM.YYYY HH:MM"
        if isinstance(conf, (int, float)):
            ws2.cell(row=r, column=9).number_format = "0%"
        if not it:
            ws2.cell(row=r, column=2).font = Font(italic=True, color="80868B")
    ws2.freeze_panes = "A2"
    if ordered:
        ws2.auto_filter.ref = f"A1:{get_column_letter(len(headers2))}{len(ordered) + 1}"

    wb.properties.title = f"Исполнение ТЗ — {project.get('name') or ''}".strip(" —")
    wb.properties.creator = "ФотоОтчёт"
    os.makedirs(os.path.dirname(os.path.abspath(out_path)) or ".", exist_ok=True)
    wb.save(out_path)
    return out_path

