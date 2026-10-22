"""Отчёт Word (.docx): титул со сводкой, таблица исполнения, пункты с фото сеткой."""

from __future__ import annotations

import datetime as _dt
import io
import os

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .common import (STATUS_COLORS, STATUS_LABELS, STATUS_TINTS, ImageCache, PhotoIndex, fmt_date, item_label,
                     norm_options, only_status_text, period_text, photo_caption, qty_text, same_text, section_info,
                     select_items, status_key, summary)

FONT = "Arial"           # есть на Windows и macOS, кириллица поддерживается
PAGE_W, MARGIN = 21.0, 1.8
USABLE_W = PAGE_W - 2 * MARGIN


def _rgb(hex6: str) -> RGBColor:
    return RGBColor.from_string(hex6.upper())


def _set_font(style_or_run, name: str = FONT, size: float | None = None, bold: bool | None = None,
              color: str | None = None) -> None:
    font = style_or_run.font
    font.name = name
    if size:
        font.size = Pt(size)
    if bold is not None:
        font.bold = bold
    if color:
        font.color.rgb = _rgb(color)
    rpr = style_or_run.element.get_or_add_rPr() if hasattr(style_or_run, "element") else style_or_run._r.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]


def _shade(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def _cell_margins(table, top=40, bottom=40, left=70, right=70) -> None:
    tblpr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblpr.append(mar)


def _row_flags(row, header: bool = False, cant_split: bool = True) -> None:
    trpr = row._tr.get_or_add_trPr()
    if cant_split:
        trpr.append(OxmlElement("w:cantSplit"))
    if header:
        trpr.append(OxmlElement("w:tblHeader"))


def _widths(table, widths_cm: list[float]) -> None:
    table.autofit = False
    grid = table._tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
        if i < len(widths_cm):
            gc.set(qn("w:w"), str(int(widths_cm[i] * 567)))
    for row in table.rows:
        for i, w in enumerate(widths_cm):
            if i < len(row.cells):
                row.cells[i].width = Cm(w)


def _para(container, text: str = "", size: float = 10, bold: bool = False, color: str | None = None,
          italic: bool = False, align=None, space_after: float = 3, keep_next: bool = False):
    p = container.add_paragraph()
    if text:
        r = p.add_run(text)
        _set_font(r, size=size, bold=bold, color=color)
        r.italic = italic
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(0)
    pf.keep_with_next = keep_next
    return p


def _run(p, text: str, size: float = 10, bold: bool = False, color: str | None = None, italic: bool = False):
    r = p.add_run(text)
    _set_font(r, size=size, bold=bold, color=color)
    r.italic = italic
    return r


def _field(p, instr: str, size: float = 8, color: str = "80868B") -> None:
    """Поле Word (PAGE, NUMPAGES) — номер страницы пересчитывается при открытии."""
    for kind, text in (("begin", None), ("instr", instr), ("separate", None), ("text", "1"), ("end", None)):
        r = p.add_run()
        _set_font(r, size=size, color=color)
        if kind in ("begin", "separate", "end"):
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            r._r.append(fc)
        elif kind == "instr":
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = f" {text} "
            r._r.append(it)
        else:
            r.text = text


def _setup(doc: Document, project: dict) -> None:
    for sec in doc.sections:
        sec.orientation = WD_ORIENT.PORTRAIT
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(MARGIN)
        sec.top_margin = Cm(1.6)
        sec.bottom_margin = Cm(1.6)
        sec.footer_distance = Cm(0.8)
        fp = sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph()
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        name = (project.get("name") or "Фотоотчёт").strip()
        _run(fp, f"{name}  ·  стр. ", size=8, color="80868B")
        _field(fp, "PAGE")
        _run(fp, " из ", size=8, color="80868B")
        _field(fp, "NUMPAGES")
    normal = doc.styles["Normal"]
    _set_font(normal, size=10)
    normal.paragraph_format.space_after = Pt(3)
    for name, size, color in (("Heading 1", 15, "202124"), ("Heading 2", 12, "202124"), ("Heading 3", 11, "3C4043"),
                              ("Title", 22, "202124")):
        st = doc.styles[name]
        _set_font(st, size=size, bold=True, color=color)
        st.paragraph_format.space_before = Pt(10 if name != "Heading 3" else 8)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.keep_with_next = True


def _heading(doc, text: str, level: int):
    h = doc.add_heading(level=level)
    r = h.add_run(text)
    _set_font(r, size={1: 15, 2: 12, 3: 11}.get(level, 11), bold=True, color="202124")
    return h


def _status_marker(p, key: str, size: float = 10) -> None:
    _run(p, "● ", size=size, color=STATUS_COLORS[key])
    _run(p, STATUS_LABELS[key], size=size, bold=True, color=STATUS_COLORS[key])


# ---------------------------------------------------------------------------
# Разделы отчёта
# ---------------------------------------------------------------------------

def _title_page(doc, project: dict, s: dict, opts: dict) -> None:
    _para(doc, "", space_after=40)
    _para(doc, "ФОТООТЧЁТ О ВЫПОЛНЕНИИ РАБОТ", size=11, bold=True, color="1A73E8", space_after=8)
    _para(doc, project.get("name") or "Проект", size=24, bold=True, color="202124", space_after=6)
    tz_title = ((project.get("tz") or {}).get("title") or "").strip()
    if tz_title:
        _para(doc, f"по документу: {tz_title}", size=11, color="3C4043", space_after=18)
    rows = []
    if project.get("customer"):
        rows.append(("Заказчик", project["customer"]))
    if project.get("contractor"):
        rows.append(("Исполнитель", project["contractor"]))
    per = period_text(s)
    if per:
        rows.append(("Период съёмки", per))
    rows.append(("Дата отчёта", _dt.date.today().strftime("%d.%m.%Y")))
    if only_status_text(opts):
        rows.append(("В отчёте пункты", only_status_text(opts)))
    t = doc.add_table(rows=len(rows), cols=2)
    _cell_margins(t)
    for i, (k, v) in enumerate(rows):
        c0, c1 = t.rows[i].cells
        c0.paragraphs[0].text = ""
        _run(c0.paragraphs[0], k, size=10, color="5F6368")
        _run(c1.paragraphs[0], str(v), size=10, bold=True)
    _widths(t, [4.0, USABLE_W - 4.0])
    _para(doc, "", space_after=14)

    _para(doc, "Сводка", size=13, bold=True, space_after=6)
    st_rows = [("Пунктов к исполнению", str(s["total"]), None)]
    for key in ("done", "provided", "executed", "todo", "na"):
        if key == "na" and not s["counts"]["na"]:
            continue
        st_rows.append((STATUS_LABELS[key], str(s["counts"][key]), key))
    if s["sections"]:
        st_rows.append(("Разделов ТЗ (заголовков)", str(s["sections"]), None))
    st_rows.append(("Готовность", f"{round(s['progress'] * 100)} %", None))
    st_rows.append(("Фото всего", str(s["photos"]), None))
    st_rows.append(("  в пунктах ТЗ / без пункта", f"{s['photos_in_items']} / {s['photos_unsorted']}", None))
    t = doc.add_table(rows=len(st_rows), cols=2)
    t.style = doc.styles["Table Grid"]
    _cell_margins(t, 60, 60, 100, 100)
    for i, (label, value, key) in enumerate(st_rows):
        c0, c1 = t.rows[i].cells
        if key:
            _status_marker(c0.paragraphs[0], key)
            _shade(c0, STATUS_TINTS[key])
            _shade(c1, STATUS_TINTS[key])
        else:
            _run(c0.paragraphs[0], label, bold=(i == 0))
        _run(c1.paragraphs[0], value, bold=True)
        c1.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _widths(t, [8.0, 3.0])


def _exec_table(doc, items: list[dict], idx: PhotoIndex, sections: dict) -> None:
    _heading(doc, "Таблица исполнения", 1)
    cols = [1.4, USABLE_W - 1.4 - 2.2 - 3.0 - 2.2 - 1.2, 2.2, 3.0, 2.2, 1.2]
    t = doc.add_table(rows=1, cols=6)
    t.style = doc.styles["Table Grid"]
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _cell_margins(t)
    for c, h in zip(t.rows[0].cells, ("№", "Пункт ТЗ", "Объём", "Статус", "Дата", "Фото")):
        _run(c.paragraphs[0], h, size=9, bold=True, color="FFFFFF")
        _shade(c, "37474F")
    _row_flags(t.rows[0], header=True)
    for it in items:
        key = status_key(it)
        row = t.add_row()
        _row_flags(row)
        cells = row.cells
        level = int(it.get("level") or 1)
        sec = sections.get(it.get("id"))
        if sec is not None:
            _run(cells[0].paragraphs[0], it.get("num") or "", size=9, bold=True)
            p = cells[1].paragraphs[0]
            p.paragraph_format.left_indent = Cm(0.35 * (level - 1))
            _run(p, it.get("title") or "", size=9, bold=True)
            if sec[1]:
                _run(cells[3].paragraphs[0], f"готово {sec[0]} из {sec[1]}", size=8, color="5F6368")
            _run(cells[5].paragraphs[0], "—", size=9)
            cells[5].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            for c in cells:
                _shade(c, "F8F9FA")
            continue
        _run(cells[0].paragraphs[0], it.get("num") or "", size=9, bold=(level == 1))
        p = cells[1].paragraphs[0]
        p.paragraph_format.left_indent = Cm(0.35 * (level - 1))
        _run(p, it.get("title") or it.get("text") or "", size=9, bold=(level == 1))
        _run(cells[2].paragraphs[0], qty_text(it), size=9)
        _run(cells[3].paragraphs[0], STATUS_LABELS[key], size=9, bold=True, color=STATUS_COLORS[key])
        _shade(cells[3], STATUS_TINTS[key])
        _run(cells[4].paragraphs[0], fmt_date(it.get("done_at")), size=9)
        n = len(idx.for_item(it))
        _run(cells[5].paragraphs[0], str(n) if n else "—", size=9)
        cells[5].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    _widths(t, cols)


def _photo_grid(doc, photos: list[dict], cache: ImageCache, per_row: int, counter: list[int]) -> int:
    """Фото сеткой per_row в ряд с подписями; возвращает число вставленных фото."""
    prepared = []
    for p in photos:
        img = cache.get(p)
        if img:
            counter[0] += 1
            prepared.append((p, img, counter[0]))
    if not prepared:
        return 0
    cell_w = USABLE_W / per_row
    img_w_max = cell_w - 0.3
    img_h_max = {1: 12.0, 2: 8.0, 3: 5.6, 4: 4.2}[per_row]
    rows = (len(prepared) + per_row - 1) // per_row
    t = doc.add_table(rows=rows, cols=per_row)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _cell_margins(t, 40, 80, 60, 60)
    for i, (photo, (data, w, h), n) in enumerate(prepared):
        cell = t.rows[i // per_row].cells[i % per_row]
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(1)
        width = min(img_w_max, img_h_max * w / max(1, h))
        p.add_run().add_picture(io.BytesIO(data), width=Cm(width))
        cap = cell.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(2)
        _run(cap, photo_caption(n, photo), size=7.5, color="5F6368")
    for row in t.rows:
        _row_flags(row)
    _widths(t, [cell_w] * per_row)
    return len(prepared)


def _item_block(doc, items: list[dict], i: int, idx: PhotoIndex, cache: ImageCache, opts: dict,
                counter: list[int], sections: dict) -> None:
    it = items[i]
    key = status_key(it)
    photos = idx.for_item(it)
    level = int(it.get("level") or 1)
    section_only = it.get("id") in sections
    _heading(doc, item_label(it) or "Пункт", 2 if level <= 1 else 3)
    if section_only:
        return
    text = (it.get("text") or "").strip()
    if text and not same_text(text, it.get("title") or ""):
        _para(doc, text, size=9, color="3C4043", space_after=3, keep_next=True)
    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(3)
    meta.paragraph_format.keep_with_next = bool(photos)
    _status_marker(meta, key)
    if it.get("done_at"):
        _run(meta, f"   ·   {fmt_date(it.get('done_at'))}", size=9, color="5F6368")
    q = qty_text(it)
    if q:
        _run(meta, f"   ·   объём: {q}", size=9, color="5F6368")
    _run(meta, f"   ·   фото: {len(photos)}", size=9, color="5F6368")
    if (it.get("comment") or "").strip():
        p = _para(doc, "", space_after=3, keep_next=bool(photos))
        _run(p, "Комментарий: ", size=9, bold=True, color="3C4043")
        _run(p, it["comment"].strip(), size=9, italic=True, color="3C4043")
    inserted = _photo_grid(doc, photos, cache, opts["photos_per_row"], counter) if photos else 0
    if not inserted and key not in ("na",):
        _para(doc, "Фото не приложены." if not photos else "Файлы фото не найдены.", size=8.5, italic=True,
              color="9AA0A6", space_after=4)


def export_docx(project: dict, project_dir: str, out_path: str, options: dict | None = None) -> str:
    """Отчёт Word: титул со сводкой, таблица исполнения, пункты с фото; возвращает out_path."""
    opts = norm_options(options)
    items = select_items(project, opts)
    idx = PhotoIndex(project, opts["skip_duplicates"])
    cache = ImageCache(project_dir, opts["max_px"], opts["jpeg_quality"])
    s = summary(project)
    sections = section_info(project, idx)

    doc = Document()
    _setup(doc, project)
    _title_page(doc, project, s, opts)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    _exec_table(doc, items, idx, sections)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    _heading(doc, "Исполнение по пунктам ТЗ", 1)
    counter = [0]
    for i in range(len(items)):
        _item_block(doc, items, i, idx, cache, opts, counter, sections)
    if opts["include_unsorted"]:
        loose = idx.unsorted()
        if loose:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            _heading(doc, "Фото без пункта ТЗ", 1)
            _para(doc, f"Фото, которые не отнесены ни к одному пункту: {len(loose)}.", size=9, color="5F6368")
            _photo_grid(doc, loose, cache, opts["photos_per_row"], counter)

    cp = doc.core_properties
    cp.title = f"Фотоотчёт — {project.get('name') or ''}".strip(" —")
    cp.author = "ФотоОтчёт"
    cp.subject = ((project.get("tz") or {}).get("title") or "")[:250]
    os.makedirs(os.path.dirname(os.path.abspath(out_path)) or ".", exist_ok=True)
    doc.save(out_path)
    return out_path

