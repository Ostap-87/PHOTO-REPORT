from __future__ import annotations
# -*- coding: utf-8 -*-
"""Перенос «Реестра подтверждений» в проект: строки реестра становятся рабочими пунктами (папками
для фото, скриншотов, файлов и ссылок), «Требуют уточнения» — отдельным списком prj["clarify"].

Источники реестра: анализ ИИ (ai.analyze) или готовый Excel в формате реестра (import_xlsx).
При повторном построении статусы, комментарии, ссылки и вложения сохраняются у строк с тем же
номером и наименованием.
"""
import io
import re
import time

from app import store

_NUM = re.compile(r"\d+(?:\.\d+)*")


def first_num(num: str) -> str:
    m = _NUM.search(num or "")
    return m.group(0) if m else ""


def _key(num: str, name: str) -> tuple[str, str]:
    return (re.sub(r"\s+", "", (num or "").rstrip(".")).lower(), re.sub(r"\s+", " ", (name or "").strip()).lower())


def full_text(prj: dict, parsed_text: str | None = None) -> str:
    """Полный текст ТЗ для анализа: из разборщика (с подставленной нумерацией) или из пунктов."""
    if parsed_text:
        return parsed_text
    lines = []
    for it in prj.get("tz_items") or prj.get("items") or []:
        head = f"{it.get('num') or ''} {it.get('title') or ''}".strip()
        body = (it.get("text") or "").strip()
        lines.append(head if not body or body.startswith(it.get("title") or "\0") else f"{head}\n{body}")
    return "\n".join(lines)


def apply(prj: dict, result: dict, mode: str = "ai") -> dict:
    """result = {"rows": [...], "clarify": [...], "summary": {...}, "model"?, "usage"?}."""
    tz_by_num = {str(it.get("num") or "").rstrip("."): it for it in prj.get("tz_items") or prj.get("items") or []}
    old = {}
    for it in prj.get("items", []):
        if it.get("kind") == "registry":
            old[_key(it.get("num"), it.get("title"))] = it
    items = []
    for r in result.get("rows") or []:
        it = store.new_item(prj, num=r.get("num", ""), title=r.get("name", ""), text=r.get("description", ""))
        it.update(kind="registry", level=1, section=r.get("section", ""), required=r.get("required", ""),
                  note=r.get("comment", ""))
        src = tz_by_num.get(first_num(r.get("num", "")))
        it["source"] = (src or {}).get("source")
        prev = old.get(_key(it["num"], it["title"]))
        if prev:
            for k in ("status", "comment", "done_at", "hint", "photo_ids", "links", "folder", "remark"):
                it[k] = prev.get(k, it[k])
            if prev.get("pf_manual"):
                it.update(pf_title=prev.get("pf_title", ""), pf_text=prev.get("pf_text", ""), pf_manual=True)
            for p in prj.get("photos", []) + prj.get("files", []):
                if p.get("item_id") == prev["id"]:
                    p["item_id"] = it["id"]
        items.append(it)
    # вложения пунктов, которых больше нет (в т.ч. пунктов ТЗ при первом реестре), — в «Неразобранное»:
    # фото затем раскладываются заново уже по строкам реестра
    items += [it for it in prj.get("items", []) if it.get("kind") == "custom"]   # свои папки пользователя остаются
    ids = {it["id"] for it in items}
    for p in prj.get("photos", []) + prj.get("files", []):
        if p.get("item_id") and p["item_id"] not in ids:
            p.update(item_id=None, confirmed=False)
    prj["items"] = items
    store.rebuild_photo_order(prj)
    prj["clarify"] = [dict(c) for c in result.get("clarify") or []]
    prj["registry"] = {"built": time.strftime("%Y-%m-%d %H:%M:%S"), "mode": mode, "stale": False,
                       "model": result.get("model", ""), "usage": result.get("usage") or {},
                       "summary": result.get("summary") or {}, "rows": len(items), "clarify": len(prj["clarify"])}
    return prj


def drop(prj: dict) -> dict:
    """Вернуться к работе по всем пунктам ТЗ (реестр удаляется, вложения — в «Неразобранное»)."""
    custom = [it for it in prj.get("items", []) if it.get("kind") == "custom"]
    keep = {it["id"] for it in custom}
    prj["items"] = [dict(it, photo_ids=[], kind="tz") for it in prj.get("tz_items") or []] + custom
    for p in prj.get("photos", []) + prj.get("files", []):
        if p.get("item_id") not in keep:
            p.update(item_id=None, confirmed=False)
    prj.pop("registry", None)
    prj["clarify"] = []
    return prj


# ------------------------------------------------------------------ импорт готового реестра из Excel
HEAD = {
    "num": ("пункт тз", "пункт", "№ пункта", "номер пункта", "№"),
    "section": ("раздел", "направление"),
    "name": ("наименование", "название"),
    "description": ("описание", "требование", "текст требования"),
    "required": ("что конкретно должно быть приложено", "что приложить", "подтверждающий материал",
                 "подтверждение"),
    "comment": ("комментарий", "примечание"),
    "reason": ("причина неопределённости", "причина неопределенности", "причина"),
}


def _match(header: str) -> str | None:
    h = re.sub(r"\s+", " ", str(header or "")).strip().lower().replace("ё", "е")
    for field, names in HEAD.items():
        for n in names:
            if h.startswith(n.replace("ё", "е")):
                return field
    return None


def _sheet_rows(ws, fields: tuple[str, ...]) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    for hi, row in enumerate(rows[:10]):
        cols = {}
        for ci, v in enumerate(row):
            f = _match(v)
            if f and f not in cols:
                cols[f] = ci
        if "num" in cols and ("required" in cols or "reason" in cols or "name" in cols):
            out = []
            for r in rows[hi + 1:]:
                rec = {f: str(r[cols[f]]).strip() if f in cols and cols[f] < len(r) and r[cols[f]] is not None else ""
                       for f in fields}
                if any(rec.values()):
                    out.append(rec)
            return out
    return []


def import_xlsx(data: bytes) -> dict:
    """Готовый реестр (листы «Реестр подтверждений» и «Требуют уточнения») -> result для apply()."""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ValueError("Не удалось открыть файл Excel реестра") from exc
    main = next((ws for ws in wb.worksheets if "реестр" in ws.title.lower()), wb.worksheets[0])
    rows = _sheet_rows(main, ("num", "section", "name", "description", "required", "comment"))
    if not rows:
        raise ValueError("В файле не найдена таблица реестра: нужны столбцы «Пункт ТЗ» и «Что конкретно должно "
                         "быть приложено» (как в «Реестре подтверждений»)")
    clar_ws = next((ws for ws in wb.worksheets if "уточн" in ws.title.lower()), None)
    clarify = _sheet_rows(clar_ws, ("num", "section", "name", "description", "reason")) if clar_ws else []
    return {"rows": rows, "clarify": clarify, "summary": {"analyzed": "", "sections": [], "issues": ""},
            "model": "", "usage": {}}
