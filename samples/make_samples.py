"""Генератор образцов ТЗ для «ФотоОтчёта» (все организации, адреса и люди вымышлены).

    python3 samples/make_samples.py

Создаёт рядом с собой:
- tz_banket.docx   — организация банкета: многоуровневая автонумерация Word (1. / 1.1. / 1.1.1.) и маркеры;
- tz_montazh.docx  — монтаж СКС: заголовки, нумерация стилем «List Number», таблица работ, маркеры;
- tz_uborka.pdf    — клининг офиса: PDF с текстовым слоем (reportlab), 2 страницы;
- tz_vystavka.xlsx — застройка выставочного стенда: таблица работ на листе Excel;
- tz_forum_table.docx — форум: всё ТЗ в одной таблице «№ п/п | Наименование раздела | Требования»,
  в ячейках — номера вручную («5.1.2.») и автонумерация Word (со startOverride и стартом уровня
  с 2), маркеры «-» и маркированные списки Word, ручные переносы строк, локальный список «1.», «2.».
"""

from __future__ import annotations

import os
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt

HERE = os.path.dirname(os.path.abspath(__file__))

CUSTOMER = "ООО «Заказчик-Пример»"
CONTRACTOR = "ООО «Подрядчик-Пример»"


# ---------------------------------------------------------------------------
# Word: общие помощники
# ---------------------------------------------------------------------------

def _base_document(font: str = "Times New Roman", size: int = 12) -> Document:
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = font
    st.font.size = Pt(size)
    rpr = st.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = parse_xml(f"<w:rFonts {nsdecls('w')}/>")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), font)
    for sec in doc.sections:
        sec.left_margin = Cm(2.5)
        sec.right_margin = Cm(1.5)
        sec.top_margin = Cm(2)
        sec.bottom_margin = Cm(2)
    return doc


def _para(doc, text: str = "", bold: bool = False, align=None, size: int | None = None, style: str | None = None,
          space_after: int = 4):
    p = doc.add_paragraph(style=style)
    if text:
        r = p.add_run(text)
        r.bold = bold
        if size:
            r.font.size = Pt(size)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p


def _header_block(doc, lines: list[str]) -> None:
    for i, line in enumerate(lines):
        _para(doc, line, bold=(i == 0), align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=0)
    _para(doc, "")


def _add_numbering(doc, abstract_xml: str, num_id: int, abstract_id: int,
                   start_overrides: dict | None = None) -> None:
    """Добавляет определение списка в numbering.xml (abstractNum — перед первым w:num).

    start_overrides: {ilvl: start} — w:lvlOverride/w:startOverride у этого списка.
    """
    numbering = doc.part.numbering_part.element
    if abstract_xml:
        abstract = parse_xml(abstract_xml)
        first_num = numbering.find(qn("w:num"))
        if first_num is not None:
            first_num.addprevious(abstract)
        else:
            numbering.append(abstract)
    overrides = "".join(f'<w:lvlOverride w:ilvl="{il}"><w:startOverride w:val="{st}"/></w:lvlOverride>'
                        for il, st in (start_overrides or {}).items())
    numbering.append(parse_xml(
        f'<w:num {nsdecls("w")} w:numId="{num_id}"><w:abstractNumId w:val="{abstract_id}"/>{overrides}</w:num>'))


def _set_num(p, num_id: int, ilvl: int) -> None:
    numpr = p._p.get_or_add_pPr().get_or_add_numPr()
    numpr.get_or_add_ilvl().val = ilvl
    numpr.get_or_add_numId().val = num_id


MULTILEVEL = f"""<w:abstractNum {nsdecls('w')} w:abstractNumId="90">
  <w:multiLevelType w:val="multilevel"/>
  <w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/>
    <w:lvlJc w:val="left"/><w:pPr><w:ind w:left="360" w:hanging="360"/></w:pPr><w:rPr><w:b/></w:rPr></w:lvl>
  <w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1.%2."/>
    <w:lvlJc w:val="left"/><w:pPr><w:ind w:left="792" w:hanging="432"/></w:pPr></w:lvl>
  <w:lvl w:ilvl="2"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1.%2.%3."/>
    <w:lvlJc w:val="left"/><w:pPr><w:ind w:left="1224" w:hanging="504"/></w:pPr></w:lvl>
</w:abstractNum>"""

BULLETS = f"""<w:abstractNum {nsdecls('w')} w:abstractNumId="91">
  <w:multiLevelType w:val="hybridMultilevel"/>
  <w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:val="–"/>
    <w:lvlJc w:val="left"/><w:pPr><w:ind w:left="1584" w:hanging="360"/></w:pPr></w:lvl>
</w:abstractNum>"""


# ---------------------------------------------------------------------------
# 1. Банкет (автонумерация Word)
# ---------------------------------------------------------------------------

BANKET = [
    ("Мебель и расстановка", [
        "Круглые банкетные столы диаметром 1,8 м — 12 шт. Расстановка по схеме Заказчика, "
        "расстояние между столами не менее 1,5 м.",
        "Банкетные стулья «Кьявари» белого цвета с мягкими подушками — 120 шт.",
        "Президиум: прямоугольные столы 1,8×0,8 м — 4 шт., установленные в одну линию у сцены.",
    ]),
    ("Текстиль и декор", [
        "Скатерти в пол белые для круглых столов — 12 шт.",
        "Чехлы на стулья с бантами цвета «шампань» — 120 шт.",
        "Салфетки льняные — 130 шт.",
        "Цветочные композиции на столы (низкие, высота не более 30 см) — 13 шт.",
    ]),
    ("Сцена", [
        "Подиум сценический 6×4 м, высота 0,4 м, с покрытием ковролином чёрного цвета.",
        "Лестница на сцену с перилами — 2 шт.",
        "Задник-баннер 6×3 м с логотипом мероприятия (макет предоставляет Заказчик).",
    ]),
    ("Звуковое оборудование", [
        "Акустическая система общей мощностью не менее 4 кВт с микшерным пультом.",
        ("Радиомикрофоны — 4 шт.", ["Комплект запасных батарей на весь период мероприятия."]),
        "Звукорежиссёр на весь период мероприятия — 1 чел.",
    ]),
    ("Световое оборудование", [
        "Заливной свет сцены (светодиодные прожекторы) — 8 шт.",
        "Архитектурная подсветка зала в цветах мероприятия.",
        "Пульт управления светом и светооператор.",
    ]),
    ("Фотозона", [
        "Пресс-волл 3×2 м с печатью по макету Заказчика.",
        "Арка из живых цветов у входа в зал.",
        "Направленный свет на фотозону — 2 шт.",
    ]),
    ("Кейтеринг", [
        ("Банкетное меню из расчёта 1 500 г на гостя, 120 персон:", None,
         ["холодные закуски и салаты;", "горячее блюдо на выбор гостя;", "десерт и торт."]),
        "Официанты — 10 чел. в униформе единого образца.",
        "Welcome-зона с напитками за 30 минут до начала банкета.",
        "Посуда, приборы и стекло из расчёта на 130 персон.",
    ]),
    ("Общие требования", [
        "Монтаж оборудования завершить не позднее 15:00 дня мероприятия.",
        "Демонтаж и вывоз оборудования — в течение 3 часов после окончания банкета.",
        "Фотоотчёт о выполнении каждого пункта предоставляется Заказчику в течение 2 рабочих дней.",
    ]),
]


def make_banket(path: str) -> None:
    doc = _base_document()
    _add_numbering(doc, MULTILEVEL, 90, 90)
    _add_numbering(doc, BULLETS, 91, 91)
    _header_block(doc, ["УТВЕРЖДАЮ", "Генеральный директор", CUSTOMER, "__________ / А.А. Образцов",
                        "«___» __________ 2026 г."])
    _para(doc, "ТЕХНИЧЕСКОЕ ЗАДАНИЕ", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=14, space_after=0)
    _para(doc, "на организацию торжественного банкета на 120 гостей", align=WD_ALIGN_PARAGRAPH.CENTER)
    _para(doc, "")
    _para(doc, "Дата проведения: 14 ноября 2026 г., с 18:00 до 23:00.")
    _para(doc, "Место проведения: банкетный зал, г. Условный, ул. Примерная, д. 1.")
    _para(doc, "Количество гостей: 120 человек.")
    _para(doc, "")
    for section, entries in BANKET:
        p = _para(doc, section, bold=True)
        _set_num(p, 90, 0)
        for entry in entries:
            if isinstance(entry, str):
                text, subs, bullets = entry, None, None
            else:
                text, subs = entry[0], entry[1]
                bullets = entry[2] if len(entry) > 2 else None
            p = _para(doc, text)
            _set_num(p, 90, 1)
            for s in subs or []:
                sp = _para(doc, s)
                _set_num(sp, 90, 2)
            for b in bullets or []:
                bp = _para(doc, b, space_after=0)
                _set_num(bp, 91, 0)
    _para(doc, "")
    _para(doc, f"Заказчик: {CUSTOMER}  __________ / А.А. Образцов")
    _para(doc, f"Исполнитель: {CONTRACTOR}  __________ / Б.Б. Примеров")
    doc.core_properties.title = "ТЗ на организацию банкета (образец)"
    doc.core_properties.author = "ФотоОтчёт — образец"
    doc.save(path)


# ---------------------------------------------------------------------------
# 2. Монтаж (таблица работ)
# ---------------------------------------------------------------------------

MONTAZH = [
    ("Кабельные трассы", [
        ("Прокладка кабеля витая пара U/UTP кат. 6 в гофрированной трубе", "м", "1 850"),
        ("Прокладка кабеля в кабель-канале 100×60 мм", "м", "240"),
        ("Монтаж кабель-канала 100×60 мм с крышкой", "м", "120"),
        ("Устройство отверстий в стенах диаметром до 50 мм", "шт.", "36"),
    ]),
    ("Лотки", [
        ("Монтаж металлического перфорированного лотка 200×50 мм за подвесным потолком", "м", "160"),
        ("Монтаж кронштейнов и шпилек для лотка", "шт.", "110"),
        ("Заземление кабельных лотков", "компл.", "1"),
    ]),
    ("Серверная", [
        ("Монтаж телекоммуникационного шкафа 42U 600×1000 мм", "шт.", "2"),
        ("Монтаж и расшивка патч-панелей 24 порта кат. 6", "шт.", "8"),
        ("Монтаж горизонтальных организаторов кабеля", "шт.", "8"),
        ("Тестирование кабельных линий с выдачей протоколов", "линия", "96"),
        ("Монтаж сплит-системы для серверной 3,5 кВт", "компл.", "1"),
    ]),
    ("Электрощиты", [
        ("Монтаж распределительного щита ЩР-1 на 24 модуля", "шт.", "1"),
        ("Прокладка кабеля ВВГнг-LS 3×2,5 от ЩР-1 до серверной", "м", "45"),
        ("Монтаж автоматических выключателей", "шт.", "12"),
        ("Монтаж розеток 220 В в серверной", "шт.", "6"),
    ]),
]


def _shade(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    tcpr.append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{fill}"/>'))


def make_montazh(path: str) -> None:
    doc = _base_document()
    _para(doc, "Приложение № 1", align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=0)
    _para(doc, "к договору подряда № 00-00/2026", align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=0)
    _para(doc, "от «___» __________ 2026 г.", align=WD_ALIGN_PARAGRAPH.RIGHT)
    _para(doc, "ТЕХНИЧЕСКОЕ ЗАДАНИЕ", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=14, space_after=0)
    _para(doc, "на монтаж структурированной кабельной системы, серверной и электрощитов",
          align=WD_ALIGN_PARAGRAPH.CENTER)
    _para(doc, "Объект: офис, г. Условный, ул. Примерная, д. 2, 3 этаж.")
    _para(doc, "Срок выполнения работ: 20 рабочих дней с даты подписания договора.")

    doc.add_heading("Общие требования", level=1)
    for t in ["Работы выполняются в нерабочее время (с 19:00 до 08:00) и в выходные дни.",
              "Исполнитель обеспечивает вывоз строительного мусора ежедневно.",
              "Материалы и оборудование согласовываются с Заказчиком до закупки."]:
        _para(doc, t, style="List Number")

    doc.add_heading("Перечень и объём работ", level=1)
    rows_total = 2 + sum(1 + len(w) for _, w in MONTAZH)
    table = doc.add_table(rows=rows_total, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [Cm(1.4), Cm(10.6), Cm(2.0), Cm(2.0)]
    header = ["№ п/п", "Наименование работ", "Ед. изм.", "Кол-во"]
    for i, h in enumerate(header):
        c = table.cell(0, i)
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        _shade(c, "D9E2F3")
    for i in range(4):
        table.cell(1, i).text = str(i + 1)
    r_i = 2
    n = 0
    for section, works in MONTAZH:
        a = table.cell(r_i, 0)
        b = table.cell(r_i, 3)
        m = a.merge(b)
        m.text = ""
        run = m.paragraphs[0].add_run(section)
        run.bold = True
        _shade(m, "F2F2F2")
        r_i += 1
        for name, unit, qty in works:
            n += 1
            vals = [str(n), name, unit, qty]
            for i, v in enumerate(vals):
                table.cell(r_i, i).text = v
            r_i += 1
    for row in table.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = w

    doc.add_heading("Требования к выполнению работ", level=1)
    for t in ["Кабели маркируются с двух сторон несмываемыми бирками.",
              "Запас кабеля в шкафу — не менее 1,5 м на каждую линию.",
              "Отверстия после прокладки заделываются негорючим материалом."]:
        _para(doc, t, style="List Bullet")

    doc.add_heading("Сдача-приёмка работ", level=1)
    _para(doc, "1) Исполнительная документация: кабельный журнал, схемы трасс, протоколы тестирования.")
    _para(doc, "2) Фотоотчёт по каждому пункту перечня работ — до закрытия потолков и после.")

    _para(doc, "")
    sig = doc.add_table(rows=1, cols=2)
    sig.cell(0, 0).text = f"Заказчик:\n{CUSTOMER}\n__________ / В.В. Условный"
    sig.cell(0, 1).text = f"Исполнитель:\n{CONTRACTOR}\n__________ / Г.Г. Примерный"
    doc.core_properties.title = "ТЗ на монтаж СКС (образец)"
    doc.core_properties.author = "ФотоОтчёт — образец"
    doc.save(path)


# ---------------------------------------------------------------------------
# 3. Уборка (PDF)
# ---------------------------------------------------------------------------

def _find_font() -> tuple[str, str]:
    cands = [
        ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
        ("/System/Library/Fonts/Supplemental/Arial.ttf", "/System/Library/Fonts/Supplemental/Arial Bold.ttf"),
        ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf"),
    ]
    for reg, bold in cands:
        if os.path.exists(reg):
            return reg, bold if os.path.exists(bold) else reg
    raise SystemExit("Не найден шрифт с кириллицей для PDF")


def make_uborka(path: str) -> None:
    from reportlab.lib.enums import TA_CENTER, TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

    reg, bold = _find_font()
    pdfmetrics.registerFont(TTFont("SampleSans", reg))
    pdfmetrics.registerFont(TTFont("SampleSans-Bold", bold))
    base = ParagraphStyle("base", fontName="SampleSans", fontSize=11, leading=15)
    right = ParagraphStyle("right", parent=base, alignment=TA_RIGHT, leading=14)
    right_b = ParagraphStyle("rightb", parent=right, fontName="SampleSans-Bold")
    title = ParagraphStyle("title", parent=base, fontName="SampleSans-Bold", fontSize=14, alignment=TA_CENTER,
                           leading=18)
    subtitle = ParagraphStyle("subtitle", parent=base, alignment=TA_CENTER, spaceAfter=10)
    sec = ParagraphStyle("sec", parent=base, fontName="SampleSans-Bold", spaceBefore=8, spaceAfter=2)
    item = ParagraphStyle("item", parent=base, leftIndent=12 * mm, firstLineIndent=-10 * mm, spaceAfter=2)
    bullet = ParagraphStyle("bullet", parent=base, leftIndent=12 * mm, firstLineIndent=-5 * mm, spaceAfter=2)

    s = []
    for i, line in enumerate(["УТВЕРЖДАЮ", f"Директор {CUSTOMER}", "__________ / В.В. Условный",
                              "«___» __________ 2026 г."]):
        s.append(Paragraph(line, right_b if i == 0 else right))
    s.append(Spacer(1, 10 * mm))
    s.append(Paragraph("ТЕХНИЧЕСКОЕ ЗАДАНИЕ", title))
    s.append(Paragraph("на оказание услуг по комплексной уборке офисных помещений", subtitle))
    s.append(Paragraph("Объект: офисные помещения общей площадью 850 м², г. Условный, ул. Примерная, д. 3. "
                       "Режим работы офиса: понедельник — пятница, с 9:00 до 19:00.", base))
    s.append(Spacer(1, 4 * mm))

    def section(t: str) -> None:
        s.append(Paragraph(t, sec))

    def it(num: str, t: str) -> None:
        s.append(Paragraph(f"{num}&nbsp;&nbsp;{t}", item))

    section("1. Ежедневная уборка")
    it("1.1.", "Влажная уборка полов во всех помещениях — 850 м².")
    it("1.2.", "Удаление пыли с рабочих столов, подоконников, оргтехники и офисной мебели.")
    it("1.3.", "Вынос мусора из корзин с заменой пакетов — 60 шт.")
    it("1.4.", "Уборка санузлов с дезинфекцией сантехники — 6 шт.")
    it("1.5.", "Пополнение расходных материалов: жидкое мыло, бумажные полотенца, туалетная бумага.")
    section("2. Еженедельная уборка")
    it("2.1.", "Мытьё стеклянных перегородок с двух сторон — 120 м².")
    it("2.2.", "Чистка ковровых покрытий пылесосом с моющей функцией — 300 м².")
    it("2.3.", "Протирка дверей, дверных ручек и выключателей — 40 шт.")
    section("3. Генеральная уборка (один раз в месяц)")
    it("3.1.", "Мытьё окон изнутри и снаружи, включая рамы, откосы и подоконники, с применением "
               "профессиональных средств, не оставляющих разводов; работы на высоте выполняются "
               "с соблюдением требований охраны труда — 210 м².")
    it("3.2.", "Химчистка мягкой мебели — 25 шт.")
    it("3.3.", "Удаление пыли с осветительных приборов и вентиляционных решёток.")
    s.append(PageBreak())
    section("4. Требования к исполнителю")
    for t in ["Персонал в униформе с логотипом исполнителя.",
              "Использование сертифицированных моющих и дезинфицирующих средств.",
              "Ведение журнала уборки с отметками о выполнении."]:
        s.append(Paragraph(f"•&nbsp;&nbsp;{t}", bullet))
    section("5. Отчётность")
    s.append(Paragraph("а)&nbsp;&nbsp;ежедневный фотоотчёт по санузлам;", bullet))
    s.append(Paragraph("б)&nbsp;&nbsp;ежемесячный акт оказанных услуг с фотографиями.", bullet))
    s.append(Spacer(1, 12 * mm))
    s.append(Paragraph(f"Заказчик: {CUSTOMER}&nbsp;&nbsp;__________ / В.В. Условный", base))
    s.append(Spacer(1, 4 * mm))
    s.append(Paragraph(f"Исполнитель: {CONTRACTOR}&nbsp;&nbsp;__________ / Г.Г. Примерный", base))

    def footer(canvas, doc_) -> None:
        canvas.saveState()
        canvas.setFont("SampleSans", 9)
        canvas.drawCentredString(A4[0] / 2, 12 * mm, f"Страница {doc_.page}")
        canvas.restoreState()

    pdf = SimpleDocTemplate(path, pagesize=A4, leftMargin=25 * mm, rightMargin=15 * mm, topMargin=20 * mm,
                            bottomMargin=20 * mm, title="ТЗ на уборку офиса (образец)", author="ФотоОтчёт — образец")
    pdf.build(s, onFirstPage=footer, onLaterPages=footer)


# ---------------------------------------------------------------------------
# 4. Выставка (Excel)
# ---------------------------------------------------------------------------

VYSTAVKA = [
    (1, "Конструкция стенда", [
        ("1.1", "Подиум высотой 100 мм с ковровым покрытием", "серый ковролин, алюминиевая окантовка", "м²", 36),
        ("1.2", "Стены из системного конструктива высотой 2,5 м", "белые панели", "п.м.", 12),
        ("1.3", "Подвесная конструкция с логотипом над стендом", "круг диаметром 3 м", "шт.", 1),
        ("1.4", "Подсобное помещение 2×1 м с дверью и замком", "", "шт.", 1),
    ]),
    (2, "Мебель", [
        ("2.1", "Стойка ресепшн с подсветкой", "брендирование по макету Заказчика", "шт.", 1),
        ("2.2", "Барные стулья", "белые", "шт.", 4),
        ("2.3", "Стол переговорный круглый", "диаметр 0,8 м", "шт.", 2),
        ("2.4", "Стулья для переговорной зоны", "", "шт.", 8),
        ("2.5", "Витрина стеклянная с подсветкой", "", "шт.", 2),
    ]),
    (3, "Электрика и свет", [
        ("3.1", "Подключение электроэнергии 10 кВт", "с щитком и УЗО", "компл.", 1),
        ("3.2", "Светильники-споты на штанге", "", "шт.", 12),
        ("3.3", "Розетки", "", "шт.", 6),
    ]),
    (4, "Графика и мультимедиа", [
        ("4.1", "Печать и монтаж фриза с названием компании", "", "п.м.", 12),
        ("4.2", "Широкоформатная печать на стены", "плёнка с ламинацией", "м²", 24),
        ("4.3", "Светодиодный экран 2×1,5 м с подключением", "", "шт.", 1),
    ]),
    (5, "Сервис", [
        ("5.1", "Ежедневная уборка стенда", "до открытия выставки", "дн.", 4),
        ("5.2", "Демонтаж и вывоз конструкций", "после окончания выставки", "усл.", 1),
    ]),
]


def make_vystavka(path: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "ТЗ стенд"
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws["A1"] = "Техническое задание на застройку выставочного стенда"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:E1")
    ws["A2"] = "Выставка «Пример-Экспо 2026», павильон 1, стенд B-12, площадь 36 м² (6×6 м), открыт с двух сторон."
    ws.merge_cells("A2:E2")
    ws["A3"] = f"Заказчик: {CUSTOMER}"
    ws.merge_cells("A3:E3")
    header = ["№ п/п", "Наименование", "Описание / требования", "Ед. изм.", "Кол-во"]
    ws.append([])
    ws.append(header)
    for c in ws[5]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="D9E2F3")
        c.border = border
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    count = 0
    for num, section, rows in VYSTAVKA:
        ws.append([num, section, None, None, None])
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="F2F2F2")
            c.border = border
        for r in rows:
            ws.append(list(r))
            count += 1
            for c in ws[ws.max_row]:
                c.border = border
                c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.append([])
    ws.append([None, f"Итого позиций: {count}"])
    for col, w in zip("ABCDE", (8, 52, 36, 10, 9)):
        ws.column_dimensions[col].width = w
    ws2 = wb.create_sheet("Контакты")
    ws2.append(["Контактное лицо", "Телефон", "Почта"])
    ws2.append(["Образцов А.А.", "+7 (000) 000-00-00", "example@example.com"])
    wb.properties.title = "ТЗ на застройку стенда (образец)"
    wb.properties.creator = "ФотоОтчёт — образец"
    wb.save(path)


# ---------------------------------------------------------------------------
# 5. Форум (всё ТЗ в одной таблице, автонумерация внутри ячеек)
# ---------------------------------------------------------------------------

def _multilevel(aid: int, starts: tuple[int, int, int]) -> str:
    """Многоуровневый список «%1.» / «%1.%2.» / «%1.%2.%3.» с заданными стартами уровней."""
    lv = []
    for il, (text, st) in enumerate(zip(("%1.", "%1.%2.", "%1.%2.%3."), starts)):
        lv.append(f'<w:lvl w:ilvl="{il}"><w:start w:val="{st}"/><w:numFmt w:val="decimal"/>'
                  f'<w:lvlText w:val="{text}"/><w:lvlJc w:val="left"/></w:lvl>')
    return (f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{aid}"><w:multiLevelType w:val="multilevel"/>'
            + "".join(lv) + "</w:abstractNum>")


FORUM_BULLETS = f"""<w:abstractNum {nsdecls('w')} w:abstractNumId="95">
  <w:multiLevelType w:val="hybridMultilevel"/>
  <w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:val="•"/>
    <w:lvlJc w:val="left"/><w:pPr><w:ind w:left="720" w:hanging="360"/></w:pPr></w:lvl>
</w:abstractNum>"""


def make_forum(path: str) -> None:
    doc = _base_document(size=11)
    # 3.x — список, у которого уровень 0 начинается с 1, но в этом документе переопределён на 3
    _add_numbering(doc, _multilevel(92, (1, 1, 1)), 92, 92, start_overrides={0: 3, 1: 1})
    # 5.2… — уровень 0 начинается с 5, уровень 1 — сразу со 2 (продолжение ручной нумерации 5.1)
    _add_numbering(doc, _multilevel(93, (5, 2, 1)), 93, 93)
    # 6.x — тот же приём, что у 3.x, но со стартом 6
    _add_numbering(doc, _multilevel(94, (1, 1, 1)), 94, 94, start_overrides={0: 6, 1: 1})
    _add_numbering(doc, FORUM_BULLETS, 95, 95)

    _header_block(doc, ["УТВЕРЖДАЮ", "Генеральный директор", CUSTOMER, "__________ / А.А. Образцов",
                        "«___» __________ 2026 г."])
    _para(doc, "ТЕХНИЧЕСКОЕ ЗАДАНИЕ", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=14, space_after=0)
    _para(doc, "на оказание услуг по организации и проведению отраслевого форума «Пример-Форум 2026»",
          align=WD_ALIGN_PARAGRAPH.CENTER)

    # Строки: (№, раздел, [абзацы ячейки]); абзац — строка (номер набран вручную или без номера)
    # или кортеж (текст, numId, ilvl) для автонумерации; «\n» в тексте — ручной перенос строки.
    rows = [
        ("1.", "Наименование услуги", [
            "1.1. Исполнитель оказывает услуги по организации и проведению отраслевого форума для субъектов "
            "малого бизнеса (далее — Форум).",
        ]),
        ("2.", "Цель", [
            "2.1. Обеспечить участие в Форуме не менее 300 участников, в том числе:",
            "- предпринимателей;",
            "- представителей отраслевых ассоциаций.",
            "2.2. Провести не менее 6 деловых сессий.",
        ]),
        ("3.", "Место и формат оказания услуг", [
            ("Услуги оказываются в очном формате на площадке, согласованной с Заказчиком.", 92, 1),
            ("Онлайн-трансляция Форума ведётся на сайте Заказчика.", 92, 1),
        ]),
        ("4.", "Сроки оказания услуг", [
            "4.1. Срок оказания услуг — с даты заключения Договора по 30 ноября 2026 г.",
            "4.2. Дата проведения Форума — 20 ноября 2026 г.",
        ]),
        ("5.", "Требования к оказываемым услугам", [
            "5.1. Исполнитель обеспечивает площадку для проведения Форума.",
            "5.1.1. Общая площадь площадки — не менее 900 кв. м.",
            "5.1.2. Оснащение главного зала:",
            "- не менее 300 посадочных мест;",
            "- сцена размером не менее 8×5 м;",
            "- светодиодный экран 6×3 м;",
            "- акустическая система — не менее 8 шт.",
            "Перечень оборудования согласовывается с Заказчиком.",
            "5.1.2.1. Резервное оборудование:",
            "1. Резервный ноутбук с презентациями.",
            "2. Запасные радиомикрофоны — 2 шт.",
            ("Исполнитель обеспечивает регистрацию участников.", 93, 1),
            ("Регистрация ведётся на сайте Форума.", 93, 2),
            ("На входе работает стойка регистрации с бейджами.", 93, 2),
            ("Исполнитель размещает логотип Форума:", 93, 1),
            ("на всех экранах площадки;", 95, 0),
            ("в окне онлайн-трансляции.", 95, 0),
            ("Подтверждение участия отправляется на почту.\nПовторное письмо — за сутки до Форума.", 93, 2),
            "5.4. Исполнитель организует кофе-брейк для 300 участников.",
        ]),
        ("6.", "Требования к отчёту", [
            ("Исполнитель предоставляет Заказчику отчёт, включающий:", 94, 1),
            ("Список участников Форума.", 94, 2),
            ("Скриншоты в количестве 15 штук, подтверждающие публикации о Форуме.", 94, 2),
            ("Фотографии, сгруппированные по каждому мероприятию Форума, — не менее 10 по каждой сессии.",
             94, 2),
            ("Видеозапись главной сессии Форума.", 94, 2),
            "6.1.5. Акт оказанных услуг.",
            ("Отчёт согласовывается с Заказчиком в течение 5 рабочих дней.", 94, 1),
        ]),
        ("7.", "Ответственные лица", [
            "От Заказчика: Образцов А.А., тел. +7 (000) 000-00-00.",
            "От Исполнителя: Примерный Г.Г., тел. +7 (000) 000-00-01.",
        ]),
    ]
    table = doc.add_table(rows=1 + len(rows), cols=3)
    table.style = "Table Grid"
    widths = [Cm(1.4), Cm(4.0), Cm(11.6)]
    for i, h in enumerate(("№ п/п", "Наименование раздела", "Требования к оказанию услуг по Договору")):
        c = table.cell(0, i)
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        _shade(c, "D9E2F3")
    for ri, (num, name, paras) in enumerate(rows, start=1):
        table.cell(ri, 0).paragraphs[0].add_run(num)
        table.cell(ri, 1).paragraphs[0].add_run(name)
        cell = table.cell(ri, 2)
        for k, entry in enumerate(paras):
            p = cell.paragraphs[0] if k == 0 else cell.add_paragraph()
            text, num_id, ilvl = (entry, None, None) if isinstance(entry, str) else entry
            for j, part in enumerate(text.split("\n")):
                if j:
                    p.runs[-1].add_break()
                p.add_run(part)
            if num_id is not None:
                _set_num(p, num_id, ilvl)
    for row in table.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = w
    _para(doc, "")
    _para(doc, f"Заказчик: {CUSTOMER}  __________ / А.А. Образцов")
    _para(doc, f"Исполнитель: {CONTRACTOR}  __________ / Г.Г. Примерный")
    doc.core_properties.title = "ТЗ на организацию форума (образец)"
    doc.core_properties.author = "ФотоОтчёт — образец"
    doc.save(path)


def main(out_dir: str = HERE) -> list[str]:
    os.makedirs(out_dir, exist_ok=True)
    paths = []
    for name, fn in (("tz_banket.docx", make_banket), ("tz_montazh.docx", make_montazh),
                     ("tz_uborka.pdf", make_uborka), ("tz_vystavka.xlsx", make_vystavka),
                     ("tz_forum_table.docx", make_forum)):
        p = os.path.join(out_dir, name)
        fn(p)
        paths.append(p)
    return paths


if __name__ == "__main__":
    for p in main(sys.argv[1] if len(sys.argv) > 1 else HERE):
        print(p, os.path.getsize(p), "байт")
